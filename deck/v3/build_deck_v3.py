"""
Deck v3 Generator for PaytmAI VyaparVoice (Post-Stage-2 Gate: stage2-ok).
Produces:
  - 5-Slide Lean Main Deck (>=18pt body text, <=5 real PPTX bullets/slide, full speaker notes)
  - 13-Slide Appendix (A1-A11 updated + A12 Cap Editor, Two-Step Voice & Stop-Loss + A13 Bayesian Scorecard & Margin x Fee Sensitivity)
Traced strictly to results/slide_numbers.md and results/unit_economics.md.
"""
import os
import re
import shutil
import sys
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import qrcode

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement

sys.path.insert(0, "/home/user/paytmai-vyaparvoice")
from app.database import init_and_seed_db
from app.engine import get_dashboard_overview
from app.stats import (
    compute_cap_binding_risk,
    compute_margin_fee_sensitivity_table,
    compute_pooled_pilot_power_table,
    simulate_dose_response_curve,
)

ASSET_DIR = "/home/user/scripts/v3_assets"
OUT_DIRS = [
    "/home/user/paytmai-vyaparvoice/deck/v3",
    "/home/user/deck/v3",
]
for d in [ASSET_DIR] + OUT_DIRS:
    os.makedirs(d, exist_ok=True)

PPTX_V3_OUT = os.path.join(OUT_DIRS[0], "PaytmAI_VyaparVoice_v3.pptx")
LIVE_URL = "https://paytmai-vyaparvoice.onrender.com"

NAVY_DARK = RGBColor(0x00, 0x1B, 0x4D)
NAVY      = RGBColor(0x00, 0x29, 0x70)
CYAN      = RGBColor(0x00, 0xBA, 0xF2)
CYAN_DARK = RGBColor(0x02, 0x84, 0xC7)
CYAN_SOFT = RGBColor(0xE0, 0xF7, 0xFE)
BG_LIGHT  = RGBColor(0xF5, 0xF8, 0xFC)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
INK       = RGBColor(0x0F, 0x17, 0x2A)
MUTED     = RGBColor(0x33, 0x41, 0x55)
BORDER    = RGBColor(0xCB, 0xD5, 0xE1)
EMERALD   = RGBColor(0x05, 0x96, 0x69)
EMER_SOFT = RGBColor(0xEC, 0xFD, 0xF5)
AMBER     = RGBColor(0xB4, 0x53, 0x09)
AMB_SOFT  = RGBColor(0xFF, 0xFB, 0xEB)
ROSE      = RGBColor(0xBE, 0x12, 0x3C)
ROSE_SOFT = RGBColor(0xFF, 0xF1, 0xF2)
SLATE_BG  = RGBColor(0xF1, 0xF5, 0xF9)

FONT_MAIN = "Liberation Sans"

EMOJI_RE = re.compile(
    "["
    "\U0001F300-\U0001FAFF"
    "\U00002700-\U000027BF"
    "\U00002600-\U000026FF"
    "]+"
)


def assert_clean_text(text: str) -> str:
    if "\u2014" in text or "\u2013" in text:
        raise ValueError(f"Banned dash found in text: {text!r}")
    if EMOJI_RE.search(text):
        raise ValueError(f"Banned emoji found in text: {text!r}")
    if text.lstrip().startswith("* "):
        raise ValueError(f"Literal asterisk bullet forbidden: {text!r}")
    if "EN reviewed" in text:
        raise ValueError(f"Forbidden phrase 'EN reviewed' found: {text!r}")
    if "[BUILT IN STAGE 1]" in text:
        raise ValueError(f"Use '[BUILT]' instead of '[BUILT IN STAGE 1]': {text!r}")
    return text


def apply_real_pptx_bullet(paragraph, bullet_char: str = "\u2022", indent_emu: int = 210000) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(indent_emu))
    pPr.set("indent", str(-indent_emu))
    for child in list(pPr):
        if child.tag.endswith(("buNone", "buChar", "buAutoNum")):
            pPr.remove(child)
    buChar = OxmlElement("a:buChar")
    buChar.set("char", bullet_char)
    pPr.append(buChar)


def generate_v3_assets():
    plt.rcParams["font.sans-serif"] = ["DejaVu Sans", "Liberation Sans", "Arial"]
    plt.rcParams["axes.edgecolor"] = "#CBD5E1"
    plt.rcParams["axes.linewidth"] = 0.9

    qr = qrcode.QRCode(version=2, error_correction=qrcode.constants.ERROR_CORRECT_M, box_size=10, border=2)
    qr.add_data(LIVE_URL)
    qr.make(fit=True)
    qr_img = qr.make_image(fill_color="#001B4D", back_color="#FFFFFF")
    qr_path = os.path.join(ASSET_DIR, "qr_live_demo.png")
    qr_img.save(qr_path)

    # Chart 1: Synthetic Store 30d Hourly Scans
    init_and_seed_db(force_reset=False)
    ov = get_dashboard_overview()
    hrs = ov["hourly_series"]
    labels = [h["label"] for h in hrs]
    counts = [h["txns"] for h in hrs]
    cols = ["#D97706" if h["is_dead_hour"] else "#002970" for h in hrs]

    fig, ax = plt.subplots(figsize=(5.0, 3.35), dpi=260)
    fig.patch.set_facecolor("#FFFFFF")
    ax.set_facecolor("#FFFFFF")
    ax.bar(labels, counts, color=cols, width=0.65, zorder=3)
    ax.grid(axis="y", linestyle="--", alpha=0.45, color="#CBD5E1", zorder=0)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.set_ylabel("30-Day UPI Scans", fontsize=9.5, fontweight="bold", color="#1E293B")
    ax.set_title("Synthetic Store 30d Hourly Scans (n = 917)", fontsize=10.5, fontweight="bold", color="#002970", pad=8)
    ax.tick_params(axis="x", labelsize=7.5, colors="#334155", rotation=30)
    ax.tick_params(axis="y", labelsize=8.5, colors="#334155")
    ax.set_ylim(0, 122)
    ax.annotate(
        "3-5 PM Slow Window\n(30 scans vs 101 peak = -70%)",
        xy=(8.0, 31),
        xytext=(4.0, 101),
        arrowprops=dict(arrowstyle="->", color="#B45309", lw=1.5),
        fontsize=8.2,
        fontweight="bold",
        color="#B45309",
        bbox=dict(boxstyle="round,pad=0.3", fc="#FFFBEB", ec="#D97706", lw=1),
    )
    plt.tight_layout()
    c1 = os.path.join(ASSET_DIR, "c1_hourly_synthetic.png")
    plt.savefig(c1, dpi=260)
    plt.close(fig)

    # Chart 2: Simulated ITT Holdout + Wilson 95% CIs
    fig, ax = plt.subplots(figsize=(5.0, 3.35), dpi=260)
    fig.patch.set_facecolor("#FFFFFF")
    ax.set_facecolor("#FFFFFF")
    arms = ["Silent Control\n(1/8 = 12.5%)", "Nudged ITT\n(12/32 = 37.5%)"]
    means = [12.5, 37.5]
    lower_err = [12.5 - 2.24, 37.5 - 22.93]
    upper_err = [47.09 - 12.5, 54.75 - 37.5]
    ax.bar(arms, means, color=["#64748B", "#059669"], width=0.44, zorder=3)
    ax.errorbar(arms, means, yerr=[lower_err, upper_err], fmt="none", ecolor="#0F172A", elinewidth=1.8, capsize=7, zorder=4, label="Wilson 95% CI")
    ax.grid(axis="y", linestyle="--", alpha=0.45, color="#CBD5E1", zorder=0)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.set_ylabel("48h Return Rate (%)", fontsize=9.5, fontweight="bold", color="#1E293B")
    ax.set_title("SIMULATED ITT Holdout + 95% CIs (p = 0.24)", fontsize=10.2, fontweight="bold", color="#002970", pad=8)
    ax.set_ylim(0, 69)
    ax.tick_params(axis="both", labelsize=8.8, colors="#1E293B")
    ax.text(0, 49.5, "95% CI:\n2.2%-47.1%", ha="center", fontsize=8.0, color="#334155", fontweight="bold")
    ax.text(1, 57.0, "95% CI:\n22.9%-54.8%", ha="center", fontsize=8.0, color="#065F46", fontweight="bold")
    ax.legend(loc="upper left", fontsize=8.0, frameon=True, facecolor="#F8FAFC")
    plt.tight_layout()
    c2 = os.path.join(ASSET_DIR, "c2_honest_holdout.png")
    plt.savefig(c2, dpi=260)
    plt.close(fig)

    # Chart 3: Three Revenue Cases at 5% of Q2 FY27 Base (8.20L Shops)
    fig, ax = plt.subplots(figsize=(5.0, 3.35), dpi=260)
    fig.patch.set_facecolor("#FFFFFF")
    ax.set_facecolor("#FFFFFF")
    cases = ["Case C:\nSaaS Only", "Case B:\nSaaS + Promo", "Case A:\nFull Base"]
    saas = np.array([77.7, 77.7, 77.7])
    promo = np.array([0.0, 100.4, 100.4])
    credit = np.array([0.0, 0.0, 75.8])
    x = np.arange(3)
    w = 0.48
    ax.bar(x, saas, width=w, label="Rs 79/mo SaaS", color="#002970", zorder=3)
    ax.bar(x, promo, width=w, bottom=saas, label="5% Promo Fee (Rs 102/mo)", color="#0284C7", zorder=3)
    ax.bar(x, credit, width=w, bottom=saas + promo, label="Credit Referral (Rs 77/mo)", color="#059669", zorder=3)
    tots = saas + promo + credit
    for xi, tot in zip(x, tots):
        ax.text(xi, tot + 6, f"Rs {tot:.1f} Cr", ha="center", va="bottom", fontsize=9.2, fontweight="bold", color="#0F172A")
    ax.set_xticks(x)
    ax.set_xticklabels(cases, fontsize=8.8, fontweight="bold", color="#1E293B")
    ax.set_ylabel("Modelled ARR (Rs Crore)", fontsize=9.5, fontweight="bold", color="#1E293B")
    ax.set_title("Modelled ARR at 5% of 1.64 Cr Base (8.20L Shops)", fontsize=9.9, fontweight="bold", color="#002970", pad=8)
    ax.set_ylim(0, 300)
    ax.grid(axis="y", linestyle="--", alpha=0.45, color="#CBD5E1", zorder=0)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.legend(loc="upper left", fontsize=7.6, frameon=True, facecolor="#F8FAFC")
    plt.tight_layout()
    c3 = os.path.join(ASSET_DIR, "c3_three_revenue_cases_q2.png")
    plt.savefig(c3, dpi=260)
    plt.close(fig)

    # Chart 4: Dose-Response & Cap Refusal Simulation (Appendix A11)
    dr_capped = simulate_dose_response_curve(offers_inr=[10, 15, 20, 25, 30, 40], max_budget=300)
    dr_uncapped = simulate_dose_response_curve(offers_inr=[10, 15, 20, 25, 30, 40], max_budget=2000)
    offers = [f"Rs {r['offer_rupees']}" for r in dr_capped]
    funded = [r["funded_redemptions"] for r in dr_capped]
    refused = [r["refused_redemptions"] for r in dr_capped]
    uncapped_profits = [r["net_profit_pre_fee_inr"] for r in dr_uncapped]

    fig, ax1 = plt.subplots(figsize=(5.1, 3.45), dpi=260)
    fig.patch.set_facecolor("#FFFFFF")
    ax1.set_facecolor("#FFFFFF")
    x4 = np.arange(len(offers))
    ax1.bar(x4, funded, width=0.46, color="#002970", label="Paid Redemptions (<=Rs 300 Stop)", zorder=3)
    ax1.bar(x4, refused, width=0.46, bottom=funded, color="#D97706", label="Refused After Rs 300 Stop", zorder=3)
    ax1.set_xticks(x4)
    ax1.set_xticklabels(offers, fontsize=8.5, fontweight="bold")
    ax1.set_ylabel("Returning Nudged Buyers (Count)", fontsize=9.0, fontweight="bold", color="#002970")
    ax1.set_ylim(0, 23)
    ax1.grid(axis="y", linestyle="--", alpha=0.4, color="#CBD5E1", zorder=0)
    ax1.spines["top"].set_visible(False)

    ax2 = ax1.twinx()
    ax2.plot(x4, uncapped_profits, color="#059669", marker="o", lw=2.2, label="Uncapped Net Profit Pre-Fee (Rs)", zorder=5)
    ax2.axhline(0, color="#94A3B8", lw=0.9, linestyle=":")
    ax2.set_ylabel("Uncapped Net Profit Pre-Fee (Rs)", fontsize=8.8, fontweight="bold", color="#059669")
    ax2.spines["top"].set_visible(False)
    ax1.set_title("SIMULATED Dose-Response: Sweet Spot at Rs 25", fontsize=9.8, fontweight="bold", color="#002970", pad=8)
    h1, l1 = ax1.get_legend_handles_labels()
    h2, l2 = ax2.get_legend_handles_labels()
    ax1.legend(h1 + h2, l1 + l2, loc="upper left", fontsize=7.1, frameon=True, facecolor="#F8FAFC")
    plt.tight_layout()
    c4 = os.path.join(ASSET_DIR, "c4_dose_breakeven.png")
    plt.savefig(c4, dpi=260)
    plt.close(fig)

    # Chart 5: Pooled 50-Shop Stratified Power Curve (Appendix A7)
    ptbl = compute_pooled_pilot_power_table(
        n_shops=50, control_per_shop=8, treatment_per_shop=32, p_control=0.125, lifts_pp=[4, 5, 6, 7, 8, 10]
    )
    xlifts = [f"+{int(r['lift_pp'])} pp\n({r['p_treatment']*100:.1f}%)" for r in ptbl["rows"]]
    ypwr = [r["power_pct"] for r in ptbl["rows"]]
    bar_cols = ["#64748B" if p < 80.0 else "#059669" for p in ypwr]

    fig, ax = plt.subplots(figsize=(5.0, 3.35), dpi=260)
    fig.patch.set_facecolor("#FFFFFF")
    ax.set_facecolor("#FFFFFF")
    x5 = np.arange(len(xlifts))
    ax.bar(x5, ypwr, width=0.52, color=bar_cols, zorder=3)
    ax.axhline(80.0, color="#B45309", linestyle="--", lw=1.4, zorder=4, label="80% Power Target (alpha=0.05)")
    for xi, pv in zip(x5, ypwr):
        ax.text(xi, pv + 2.2, f"{pv:.1f}%", ha="center", va="bottom", fontsize=8.2, fontweight="bold", color="#0F172A")
    ax.set_xticks(x5)
    ax.set_xticklabels(xlifts, fontsize=7.8, fontweight="bold", color="#1E293B")
    ax.set_ylabel("Two-Sided Power (%)", fontsize=9.2, fontweight="bold", color="#1E293B")
    ax.set_ylim(0, 112)
    ax.set_title("50-Shop Pooled Power (400C + 1,600T, 12.5% Base)", fontsize=9.8, fontweight="bold", color="#002970", pad=8)
    ax.grid(axis="y", linestyle="--", alpha=0.45, color="#CBD5E1", zorder=0)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.legend(loc="lower right", fontsize=7.8, frameon=True, facecolor="#F8FAFC")
    plt.tight_layout()
    c5 = os.path.join(ASSET_DIR, "c5_pooled_power.png")
    plt.savefig(c5, dpi=260)
    plt.close(fig)

    # Chart 6 (NEW Stage 2): Cap Editor Presets vs Binomial Overrun Probability (Task 9)
    presets = [150, 250, 300, 400, 500]
    overrun_pcts = [
        compute_cap_binding_risk(32, 25, cap_v, 0.375, 0.10)["prob_exceeds_cap_pct"]
        for cap_v in presets
    ]
    x6 = np.arange(len(presets))
    x6_lbls = [f"Rs {c}\n({c//25} paid)" for c in presets]
    c6_cols = ["#D97706" if p > 10.0 else "#059669" for p in overrun_pcts]

    fig, ax = plt.subplots(figsize=(5.0, 3.35), dpi=260)
    fig.patch.set_facecolor("#FFFFFF")
    ax.set_facecolor("#FFFFFF")
    ax.bar(x6, overrun_pcts, width=0.50, color=c6_cols, zorder=3)
    ax.axhline(10.0, color="#BE123C", linestyle="--", lw=1.4, zorder=4, label="10% Overrun Alert Threshold")
    for xi, pv in zip(x6, overrun_pcts):
        y_off = pv + 6.2 if 3.5 <= pv <= 10.5 else pv + 2.4
        ax.text(xi, y_off, f"{pv:.1f}%", ha="center", va="bottom", fontsize=8.4, fontweight="bold", color="#0F172A")
    ax.set_xticks(x6)
    ax.set_xticklabels(x6_lbls, fontsize=8.0, fontweight="bold", color="#1E293B")
    ax.set_ylabel("P(Redemptions > Cap | p=37.5%)", fontsize=9.0, fontweight="bold", color="#1E293B")
    ax.set_ylim(0, 114)
    ax.set_title("Cap Editor Presets vs Binomial Refusal Risk", fontsize=9.8, fontweight="bold", color="#002970", pad=8)
    ax.grid(axis="y", linestyle="--", alpha=0.45, color="#CBD5E1", zorder=0)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.legend(loc="upper right", fontsize=7.8, frameon=True, facecolor="#F8FAFC")
    plt.tight_layout()
    c6 = os.path.join(ASSET_DIR, "c6_cap_overrun_presets.png")
    plt.savefig(c6, dpi=260)
    plt.close(fig)

    # Chart 7 (NEW Stage 2): 6x4 Margin x Paytm Fee Sensitivity Heatmap (Task 11)
    sens = compute_margin_fee_sensitivity_table(1740.0, 300.0, 2610.0)
    margins = [f"{int(m)}%" for m in sens["margins_pct"]]
    fees = [f"{int(f)}% Fee" for f in sens["fees_pct"]]
    mat = np.array([[cell["net_profit_inr"] for cell in row["cells"]] for row in sens["rows"]])

    fig, ax = plt.subplots(figsize=(5.1, 3.45), dpi=260)
    fig.patch.set_facecolor("#FFFFFF")
    im = ax.imshow(mat, cmap="RdYlGn", aspect="auto", vmin=-220, vmax=360)
    ax.set_xticks(np.arange(len(fees)))
    ax.set_yticks(np.arange(len(margins)))
    ax.set_xticklabels(fees, fontsize=8.5, fontweight="bold", color="#1E293B")
    ax.set_yticklabels([f"{m} Margin" for m in margins], fontsize=8.3, fontweight="bold", color="#1E293B")
    for r_i in range(len(margins)):
        for c_i in range(len(fees)):
            v = mat[r_i, c_i]
            txt = f"+Rs {v:.0f}" if v >= 0 else f"-Rs {abs(v):.0f}"
            if r_i == 3 and c_i == 2:
                txt = f"+Rs {v:.1f}*"
            fg_col = "#FFFFFF" if (v <= -140 or v >= 280) else "#0F172A"
            ax.text(c_i, r_i, txt, ha="center", va="center", fontsize=8.2, fontweight="bold", color=fg_col)
    ax.set_title("Net Profit Sensitivity (Rs): Margin x Paytm Fee (*Base)", fontsize=9.4, fontweight="bold", color="#002970", pad=8)
    plt.tight_layout()
    c7 = os.path.join(ASSET_DIR, "c7_margin_fee_heatmap.png")
    plt.savefig(c7, dpi=260)
    plt.close(fig)

    return qr_path, c1, c2, c3, c4, c5, c6, c7


def set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_box(slide, left, top, width, height, fill_rgb, border_rgb=None, border_pt=1.0, rounded=True):
    st = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(st, Inches(left), Inches(top), Inches(width), Inches(height))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill_rgb
    if border_rgb:
        shp.line.color.rgb = border_rgb
        shp.line.width = Pt(border_pt)
    else:
        shp.line.fill.background()
    if rounded and len(shp.adjustments) > 0:
        shp.adjustments[0] = 0.03
    return shp


def add_text(slide, left, top, width, height, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, italic=False):
    assert_clean_text(text)
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for i, line in enumerate(text.split("\n")):
        assert_clean_text(line)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(3)
        r = p.add_run()
        r.text = line
        r.font.name = FONT_MAIN
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
    return tb


def add_bullet_list(slide, left, top, width, height, bullets, font_pt=18.0, space_pt=12, lead_color=NAVY, body_color=INK):
    if len(bullets) > 5:
        raise ValueError(f"Exceeded 5 bullets limit: got {len(bullets)}")
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    for idx, item in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.space_after = Pt(space_pt)
        apply_real_pptx_bullet(p)
        if isinstance(item, tuple):
            lead, body = item
            assert_clean_text(lead)
            assert_clean_text(body)
            r1 = p.add_run()
            r1.text = f"{lead}: "
            r1.font.name = FONT_MAIN
            r1.font.size = Pt(font_pt)
            r1.font.bold = True
            r1.font.color.rgb = lead_color

            r2 = p.add_run()
            r2.text = body
            r2.font.name = FONT_MAIN
            r2.font.size = Pt(font_pt)
            r2.font.bold = False
            r2.font.color.rgb = body_color
        else:
            assert_clean_text(item)
            r = p.add_run()
            r.text = item
            r.font.name = FONT_MAIN
            r.font.size = Pt(font_pt)
            r.font.color.rgb = body_color
    return tb


def add_chrome(slide, tag_text: str, title: str, page_label: str, is_appendix: bool = False):
    assert_clean_text(tag_text)
    assert_clean_text(title)
    assert_clean_text(page_label)
    add_box(slide, 0, 0, 13.333, 0.08, AMBER if is_appendix else CYAN, rounded=False)

    pill_bg = AMB_SOFT if is_appendix else CYAN_SOFT
    pill_fg = AMBER if is_appendix else NAVY
    pill_w = max(2.6, len(tag_text) * 0.095 + 0.5)
    add_box(slide, 0.55, 0.24, pill_w, 0.32, pill_bg, border_rgb=pill_fg, border_pt=1.0)
    add_text(slide, 0.60, 0.27, pill_w - 0.1, 0.26, tag_text, size=9.5, bold=True, color=pill_fg, align=PP_ALIGN.CENTER)

    add_text(
        slide, 7.5, 0.26, 5.28, 0.28,
        "PaytmAI VyaparVoice | Team SemiTechnical (Kevin Joseph)",
        size=9.5, bold=True, color=MUTED, align=PP_ALIGN.RIGHT
    )
    add_text(slide, 0.55, 0.66, 12.23, 0.52, title, size=23.0, bold=True, color=INK)

    add_box(slide, 0.55, 7.08, 12.23, 0.015, BORDER, rounded=False)
    add_text(
        slide, 0.55, 7.13, 9.2, 0.26,
        "Paytm Build for India AI Hackathon (Bengaluru Track 1) | Traced to results/slide_numbers.md & unit_economics.md (Deck v3)",
        size=8.6, color=MUTED
    )
    add_text(slide, 10.4, 7.13, 2.38, 0.26, page_label, size=8.8, bold=True, color=NAVY, align=PP_ALIGN.RIGHT)


def set_notes(slide, notes_text: str):
    assert_clean_text(notes_text)
    slide.notes_slide.notes_text_frame.text = notes_text


def build_deck_v3():
    qr_path, c1, c2, c3, c4, c5, c6, c7 = generate_v3_assets()
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # MAIN 1 / 5
    m1 = prs.slides.add_slide(blank)
    set_bg(m1, BG_LIGHT)
    add_chrome(m1, "MAIN 1 / 5 : PROBLEM & OPPORTUNITY", "1.64 Cr Merchants Pay for Subscriptions; Retention Stays Manual", "Main 1 / 5 (Slide 1 / 18)")
    add_box(m1, 0.55, 1.30, 7.25, 5.55, WHITE, border_rgb=BORDER)
    add_bullet_list(
        m1, 0.72, 1.48, 6.90, 5.20,
        [
            ("Reported Paytm Scale", "1.64 Cr subscription merchants (+27L YoY), Rs 7.1L Cr GMV, Rs 2,648 Cr revenue (Reported: Paytm Q2 FY27 release, 26 Oct 2026)."),
            ("Silent Repeat Churn", "Small shops lose an estimated 25% to 35% of repeat buyers yearly without customer-level alerts (Assumption)."),
            ("Dashboard Friction", "Single-counter operators work 12+ hr shifts; SMB BI dashboards see <8% weekly usage (Assumption)."),
            ("Slow Hours & Rush Prep", "In our synthetic Bengaluru ledger, 3 to 5 PM scans fall 70% below peak while weekend rushes risk stockouts."),
        ],
        font_pt=18.0, space_pt=14
    )
    add_box(m1, 8.00, 1.30, 4.78, 3.52, WHITE, border_rgb=BORDER)
    m1.shapes.add_picture(c1, Inches(8.08), Inches(1.38), width=Inches(4.62))
    add_box(m1, 8.00, 4.98, 4.78, 1.87, NAVY_DARK, border_rgb=CYAN, border_pt=1.2)
    add_text(m1, 8.18, 5.12, 4.42, 1.60, "PaytmAI VyaparVoice (Track 1)\n8:30 PM closing-time voice retention & ops copilot for Soundbox + Paytm for Business.\nTeam SemiTechnical | Kevin Joseph", size=13.5, bold=True, color=WHITE)
    set_notes(m1, "Good morning. I am Kevin Joseph from Team SemiTechnical, presenting PaytmAI VyaparVoice for Track 1, Merchant Growth AI, anchored to Paytm's 26 October 2026 Q2 FY27 earnings release.")

    # MAIN 2 / 5
    m2 = prs.slides.add_slide(blank)
    set_bg(m2, BG_LIGHT)
    add_chrome(m2, "MAIN 2 / 5 : SOLUTION & LIVE DEMO", "One 8:30 PM Voice Action, Governed by Deterministic SQL Rules", "Main 2 / 5 (Slide 2 / 18)")
    add_box(m2, 0.55, 1.30, 7.85, 5.55, WHITE, border_rgb=BORDER)
    add_bullet_list(
        m2, 0.72, 1.44, 7.50, 5.25,
        [
            ("Plain Voice & Ergonomics", "EN author-checked, native review pending; KN/HI draft; Repeat, 0.78x Slower, Question Chips & Why? drawers."),
            ("Synthetic Store Slice", "Sri Krishna Darshini (240 payers, 1,613 scans); flags 48 regulars silent 16 to 38d."),
            ("Policy & 80/20 Split", "Filters to 40 DPDP & 14d-cooldown clear buyers: 32 nudged vs 8 silent control."),
            ("Cap Editor & Stop-Loss", "All 32 get Rs 25 pass; Cap Editor (Rs 50-2,500), 2-step voice confirm & 1-tap Stop-Loss protect pocket."),
            ("Zero LLM Math", "SQLite computes every rupee; per-visitor cloud sandboxes block paid API calls."),
        ],
        font_pt=18.0, space_pt=10
    )
    add_box(m2, 8.58, 1.30, 4.20, 5.55, NAVY_DARK, border_rgb=CYAN, border_pt=1.3)
    add_text(m2, 8.78, 1.46, 3.80, 0.36, "TRY LIVE SANDBOX DEMO", size=14.5, bold=True, color=CYAN, align=PP_ALIGN.CENTER)
    add_box(m2, 9.53, 1.92, 2.30, 2.30, WHITE, border_rgb=CYAN, border_pt=1.0)
    m2.shapes.add_picture(qr_path, Inches(9.63), Inches(2.02), width=Inches(2.10))
    add_text(m2, 8.72, 4.34, 3.92, 0.62, "https://paytmai-vyaparvoice\n.onrender.com", size=13.0, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_box(m2, 8.78, 5.06, 3.80, 1.60, NAVY, border_rgb=CYAN, border_pt=1.0)
    add_text(m2, 8.90, 5.15, 3.56, 1.44, "Cloud Safeguards Active:\nIsolated SQLite per visitor (30m TTL,\n64-session LRU, 2MB/32MB disk caps).\nZero paid API calls on public URL.", size=11.0, color=CYAN_SOFT, align=PP_ALIGN.CENTER)
    set_notes(m2, "Slide 2 highlights our built Stage 2 merchant controls: Repeat and 0.78x Slower counter playback, five plain-language Question Chips, a bounded Cap Editor with two-step voice confirmation, and an atomic 1-tap Stop-Loss.")

    # MAIN 3 / 5
    m3 = prs.slides.add_slide(blank)
    set_bg(m3, BG_LIGHT)
    add_chrome(m3, "MAIN 3 / 5 : ARCHITECTURE (BUILT VS PLANNED)", "Deterministic Core, Cap Controls & Scorecard Built; Rail Hooks Scoped", "Main 3 / 5 (Slide 3 / 18)")
    add_box(m3, 0.55, 1.30, 12.23, 4.60, WHITE, border_rgb=BORDER)
    add_bullet_list(
        m3, 0.78, 1.44, 11.75, 4.34,
        [
            ("[BUILT] ITT Stats, Cap Editor & Stop-Loss", "Fisher/Wilson/power math, Cap Editor (Rs 50-2,500), 2-step voice confirm, atomic Stop-Loss, What-If & Guardrail Lab."),
            ("[BUILT] Bayesian Scorecard & Copy Ergonomics", "Plain verdict, Beta(1,1) profit posterior, 6x4 margin-fee table, Evidence Meter, Repeat/0.78x Slower, banned-jargon lint."),
            ("[BUILT KEYLESS / PENDING LIVE] Sarvam & Caps", "Keyless Bulbul TTS + Saaras STT adapter, mock-server selftest, SHA-256 audio cache, 15%/15% (30% max) credit cap."),
            ("[PRODUCTION PLAN] Paytm Ecosystem Hooks", "Paytm QR dynamic coupon settlement, 4G Soundbox push audio, WhatsApp Cloud supplier POs, Postgres/Redis."),
        ],
        font_pt=18.0, space_pt=12
    )
    cap_tiles_m3 = [
        ("6-Rule Guardrail Stress Lab", "Blocks prompt injection, >22% discounts, unconsented blasts, and holdout removal.", EMERALD, EMER_SOFT),
        ("Atomic Cap & Stop-Loss Lock", "SQLite write locking enforces hard rupee stops and instant 1-tap campaign pause.", NAVY, CYAN_SOFT),
        ("Bounded Visitor Sandboxes", "Per-visitor SQLite files enforce 30-min TTL, 64-session LRU eviction, and 32 MB disk cap.", AMBER, AMB_SOFT),
    ]
    for gi, (gh, gb, gcol, gbg) in enumerate(cap_tiles_m3):
        gx = 0.55 + gi * 4.14
        add_box(m3, gx, 6.04, 3.94, 0.82, gbg, border_rgb=gcol, border_pt=1.1)
        add_text(m3, gx + 0.14, 6.10, 3.66, 0.24, gh, size=11.0, bold=True, color=gcol)
        add_text(m3, gx + 0.14, 6.35, 3.66, 0.44, gb, size=9.8, color=INK)
    set_notes(m3, "Slide 3 reflects our completed Stage 2 codebase: the Cap Editor with two-step voice confirmation, live interim Stop-Loss, Guardrail Red-Team Lab, What-If explorer, and Bayesian plus margin-fee sensitivity scorecard.")

    # MAIN 4 / 5
    m4 = prs.slides.add_slide(blank)
    set_bg(m4, BG_LIGHT)
    add_chrome(m4, "MAIN 4 / 5 : SIMULATED ITT PROOF & PROFIT MATH", "Intention-to-Treat Holdout, Bayesian Posterior & Net Profit", "Main 4 / 5 (Slide 4 / 18)")
    add_box(m4, 0.55, 1.30, 7.25, 4.74, WHITE, border_rgb=BORDER)
    add_bullet_list(
        m4, 0.72, 1.42, 6.92, 4.54,
        [
            ("Simulated ITT Return", "12/32 nudged (37.5%) vs 1/8 control (12.5%); Fisher p = 0.24; simulated 12 equals the Rs 300 cap (0 refused)."),
            ("Bayesian Posterior (n=40)", "Beta(1,1) prior: 88.8% P(lift>0), 67.0% P(pre-fee profit>0), 52.1% P(post-fee profit>0); proves 50-shop pool is needed."),
            ("Sales vs Subsidy", "8 net visits = Rs 1,740 sales (5.8x sales/promo); 4 organic returners absorb Rs 100 cashback."),
            ("Merchant Net Profit", "At 30% margin: Rs 522 GP (1.74x); +Rs 222 pre-fee (+74%), +Rs 91.50 after Rs 130.50 fee."),
            ("5% Fee Base", "Rs 130.50 = 5% of all Rs 2,610 redeemed bills (12 bills, including organic returners)."),
        ],
        font_pt=18.0, space_pt=9
    )
    add_box(m4, 0.55, 6.14, 7.25, 0.72, CYAN_SOFT, border_rgb=CYAN_DARK)
    add_text(
        m4, 0.70, 6.21, 6.96, 0.58,
        "Intention-to-Treat (ITT) Definition: Compares all 32 assigned nudged buyers vs all 8 assigned silent controls regardless of redemption or cap stops, preventing selection bias and keeping lift invariant if the cap binds.",
        size=11.0, bold=True, color=NAVY
    )
    add_box(m4, 8.00, 1.30, 4.78, 3.52, WHITE, border_rgb=BORDER)
    m4.shapes.add_picture(c2, Inches(8.08), Inches(1.38), width=Inches(4.62))
    add_box(m4, 8.00, 4.96, 4.78, 1.90, AMB_SOFT, border_rgb=AMBER, border_pt=1.2)
    add_text(
        m4, 8.16, 5.06, 4.46, 1.72,
        "CAP & BREAK-EVEN DISCLOSURE (SIMULATED)\nSimulated 12 returners = 12 x Rs 25 = Rs 300 hard spend stop (42.2% overrun risk at p=37.5% triggers >10% alert; suggests Rs 400 cap).\nBreak-even margin: 17.2% pre-fee | 24.7% post-fee.\nBreak-even lift (30% margin): +7.8 pp pre | +15.3 pp post.",
        size=11.0, bold=True, color=INK
    )
    set_notes(m4, "Slide 4 pairs frequentist Fisher and Wilson readouts with our built Bayesian Beta(1,1) profit posterior: on one 40-customer wave, post-fee profit probability is 52.1 percent with a 95 percent credible interval of minus Rs 19.97 to plus Rs 14.73 per nudged customer, demonstrating why our 50-shop pooled pilot is essential.")

    # MAIN 5 / 5
    m5 = prs.slides.add_slide(blank)
    set_bg(m5, BG_LIGHT)
    add_chrome(m5, "MAIN 5 / 5 : BUSINESS MODEL & POOLED PILOT", "Three Revenue Cases (8.20L Shops) & 50-Shop Pooled Pilot", "Main 5 / 5 (Slide 5 / 18)")
    add_box(m5, 0.55, 1.30, 7.25, 4.74, WHITE, border_rgb=BORDER)
    add_bullet_list(
        m5, 0.72, 1.42, 6.92, 4.54,
        [
            ("Case C (SaaS Only)", "Rs 79/mo Soundbox add-on x 8.20L shops = Rs 77.7 Cr ARR (zero promo/credit fees)."),
            ("Case B (Core Excl. Credit)", "+Rs 102/mo steady promo fee (~0.78 waves/mo, Rs 2,040 redeemed bills) = Rs 178.1 Cr ARR."),
            ("Case A (Full Base)", "+Rs 77/mo credit referral (~7% loan pen. Assumption) = Rs 253.9 Cr ARR."),
            ("Unit COGS & Margin", "Rs 48/mo cached (Rs 5 AI + Rs 43 msg; 81.4% margin) vs Rs 71.80 uncached TTS (72.2%)."),
            ("Pooled 50-Shop Pilot", "50 shops x 1 wave (~400C + ~1,600T, stratified by shop): 83.3% power at +6 pp, 97.2% at +8 pp."),
        ],
        font_pt=18.0, space_pt=9
    )
    add_box(m5, 0.55, 6.14, 7.25, 0.72, EMER_SOFT, border_rgb=EMERALD)
    add_text(
        m5, 0.70, 6.22, 6.96, 0.56,
        "Why Pool 50 Shops (Not 4 Waves/Shop): 14d cooldown & ~0.78 waves/mo replenishment preclude 4 weekly waves in 1 shop; pooling 50x1 gives 50.3%/68.8%/83.3%/92.5%/97.2%/99.8% power at +4/+5/+6/+7/+8/+10 pp.",
        size=10.8, bold=True, color=INK
    )
    add_box(m5, 8.00, 1.30, 4.78, 3.52, WHITE, border_rgb=BORDER)
    m5.shapes.add_picture(c3, Inches(8.08), Inches(1.38), width=Inches(4.62))
    add_box(m5, 8.00, 4.96, 4.78, 1.90, NAVY_DARK, border_rgb=CYAN, border_pt=1.2)
    add_text(
        m5, 8.16, 5.08, 4.46, 1.66,
        "Q2 FY27 BASE & FEE RECONCILIATION\n5% of 1.64 Cr subscription merchants (Reported: Paytm Q2 FY27, 26 Oct 2026) = 8.20L shops.\n1 backlog wave = Rs 130.50 fee; steady state (~0.78 waves/mo) = Rs 102/mo.",
        size=11.5, bold=True, color=WHITE
    )
    set_notes(m5, "Slide 5 presents our reproducible unit_economics.py outputs across 8.20 lakh shops (5 percent of Paytm's 1.64 crore Q2 FY27 merchant base) and our pooled 50-shop stratified pilot.")

    # APPENDIX A1 (Slide 6 / 18)
    a1 = prs.slides.add_slide(blank)
    set_bg(a1, NAVY_DARK)
    add_box(a1, 0, 0, 13.333, 0.10, AMBER, rounded=False)
    add_box(a1, 0.65, 0.45, 8.10, 0.42, NAVY, border_rgb=CYAN)
    add_text(a1, 0.80, 0.52, 7.80, 0.28, "APPENDIX A1 | TRACK 1: MERCHANT GROWTH AI (STAGE 2 VERIFIED - DECK V3)", size=10.0, bold=True, color=CYAN)
    add_text(a1, 0.65, 1.25, 8.2, 0.80, "PaytmAI VyaparVoice", size=40, bold=True, color=WHITE)
    add_text(a1, 0.65, 2.10, 8.2, 0.70, "Closing-Time Voice Growth, Operations & Retention Copilot", size=17.5, bold=True, color=CYAN)
    add_text(a1, 0.65, 2.92, 8.2, 0.95, "Built for Paytm's 1.64 Cr subscription-paying merchant deployments (Reported: Paytm Q2 FY27 release, 26 Oct 2026). Includes Cap Editor, 2-step voice confirmation, Stop-Loss, and Bayesian scorecard.", size=13.0, color=RGBColor(0xCB, 0xD5, 0xE1))
    cards_a1 = [
        ("01 : Voice + Teach-Back", "Repeat, 0.78x Slower, 5 Question Chips & Why? drawers"),
        ("02 : Cap Editor & Stop-Loss", "Rs 50-2,500 bounds, 2-step voice confirm & 1-tap halt"),
        ("03 : Bayesian Scorecard", "Beta(1,1) profit posterior + 6x4 margin-fee table"),
    ]
    for i, (ch, cb) in enumerate(cards_a1):
        bx = 0.65 + i * 2.75
        add_box(a1, bx, 4.15, 2.60, 1.18, NAVY, border_rgb=CYAN)
        add_text(a1, bx + 0.12, 4.25, 2.36, 0.30, ch, size=10.8, bold=True, color=WHITE)
        add_text(a1, bx + 0.12, 4.58, 2.36, 0.62, cb, size=9.6, color=RGBColor(0xBA, 0xE6, 0xFD))
    add_box(a1, 0.65, 5.65, 12.0, 1.10, RGBColor(0x00, 0x12, 0x35), border_rgb=CYAN)
    add_text(a1, 0.85, 5.80, 11.5, 0.35, "TEAM: SemiTechnical | Member: Kevin Joseph | Live: paytmai-vyaparvoice.onrender.com", size=13.5, bold=True, color=WHITE)
    add_text(a1, 0.85, 6.22, 11.5, 0.35, "Deck v3 Status: Stage 2 verified (stage2-ok); Cap Editor (A12) and Bayesian/Sensitivity Scorecard (A13) added.", size=11.5, color=CYAN)
    set_notes(a1, "Appendix A1 introduces the 13-slide Deck v3 technical appendix.")

    # APPENDIX A2 (Slide 7 / 18)
    a2 = prs.slides.add_slide(blank)
    set_bg(a2, BG_LIGHT)
    add_chrome(a2, "APPENDIX A2 : PROBLEM & Q2 FY27 SOURCES", "Problem Details, Primary Paytm Q2 FY27 Sources & Synthetic Persona", "Appendix A2 (Slide 7 / 18)", True)
    p_cards = [
        ("1. Silent Repeat Churn (25% to 35% Yearly - ASSUMPTION)", "Merchants hear isolated payment alerts without customer-level retention tracking, losing repeat buyers before cashflow drops are noticed."),
        ("2. Dashboard Friction (<8% WAU - ASSUMPTION)", "Counter-operated kirana, darshini, and pharmacy shops lack time to write prompts, export CSVs, or configure CRM rules."),
        ("3. Slow Hours & Rush Prep (SYNTHETIC STORE)", "In our 60-day synthetic Indiranagar ledger, 3 to 5 PM scans fall 70% below peak while weekend morning rushes risk staple stockouts."),
    ]
    for i, (pt, pd) in enumerate(p_cards):
        cy = 1.40 + i * 1.78
        add_box(a2, 0.55, cy, 6.45, 1.60, WHITE, border_rgb=BORDER)
        add_text(a2, 0.75, cy + 0.14, 6.05, 0.34, pt, size=13.0, bold=True, color=NAVY)
        add_text(a2, 0.75, cy + 0.52, 6.05, 0.95, pd, size=11.5, color=MUTED)
    add_box(a2, 7.20, 1.40, 5.58, 3.45, WHITE, border_rgb=BORDER)
    a2.shapes.add_picture(c1, Inches(7.54), Inches(1.46), width=Inches(4.90))
    add_box(a2, 7.20, 4.98, 5.58, 1.92, NAVY_DARK, border_rgb=CYAN)
    add_bullet_list(
        a2, 7.34, 5.08, 5.30, 1.74,
        [
            ("Reported (Paytm Q2 FY27, 26 Oct 2026)", "1.64 Cr subscription-paying merchants, Rs 7.1L Cr merchant GMV, Rs 2,648 Cr operating revenue."),
            ("Assumptions vs Synthetic", "~7% merchant loan pen. & Rs 90-100/mo Soundbox fee (Assumptions); Ramesh Rao store is synthetic (917 scans/30d)."),
        ],
        font_pt=10.6, space_pt=7, lead_color=CYAN, body_color=WHITE
    )
    set_notes(a2, "Appendix A2 cites Paytm's Q2 FY27 release and separates market assumptions from synthetic store data.")

    # APPENDIX A3 (Slide 8 / 18)
    a3 = prs.slides.add_slide(blank)
    set_bg(a3, BG_LIGHT)
    add_chrome(a3, "APPENDIX A3 : 3-PILLAR MAPPING", "Track 1 Coverage: Growth, Operations, and Customer Recovery", "Appendix A3 (Slide 8 / 18)", True)
    pillars_a3 = [
        ("PILLAR 1 : GROW BUSINESS", "Retention & Slow-Hour Copilot", [
            ("Churn Detection (Built)", "SQL flags regulars silent 16+ days (16-38d in synthetic DB)."),
            ("Slow-Hour Group (Built)", "Identifies 3-5 PM snack buyers (55 detected, 34 eligible)."),
            ("Cap & Stop-Loss (Built)", "Cap Editor (Rs 50-2,500), 2-step voice confirm & Stop-Loss."),
            ("QR Coupon Hook (Planned)", "Production hook into Paytm QR redemption rail."),
        ]),
        ("PILLAR 2 : MANAGE OPERATIONS", "Rush Prep & Credit Signals", [
            ("Rush Alert (Simulated)", "Flags weekend morning demand spikes from scan history."),
            ("Supplier PO (Simulated)", "Dispatches pre-filled WhatsApp restock (+18L milk, +12kg batter)."),
            ("Dual Ledgers (Built)", "Separates merchant cashback stop from 30% Sarvam credit cap."),
            ("Credit Signal (Planned)", "Feeds verified cashflow stability to partner NBFC funnel."),
        ]),
        ("PILLAR 3 : SERVE CUSTOMERS", "Privacy-Safe Service Recovery", [
            ("Delay Tagging (Simulated)", "Flags peak-hour queue delays and 3-star feedback tags."),
            ("Goodwill Pass (Simulated)", "1-tap bilingual apology + Rs 30 next-scan coffee voucher."),
            ("Zero PII Leakage (Built)", "No raw phone numbers or payer IDs ever leave SQLite."),
            ("DPDP & Cooldown (Built)", "Hard opt-in filter and 14-day per-payer anti-spam cooldown."),
        ]),
    ]
    for i, (ptag, phead, pbul) in enumerate(pillars_a3):
        px = 0.55 + i * 4.14
        add_box(a3, px, 1.40, 3.94, 5.45, WHITE, border_rgb=BORDER)
        add_text(a3, px + 0.18, 1.56, 3.58, 0.28, ptag, size=10.5, bold=True, color=NAVY)
        add_text(a3, px + 0.18, 1.86, 3.58, 0.36, phead, size=14.5, bold=True, color=INK)
        add_bullet_list(a3, px + 0.14, 2.32, 3.66, 4.35, pbul, font_pt=12.0, space_pt=10)
    set_notes(a3, "Appendix A3 maps built modules, simulated MVP flows, and planned Paytm hooks across all three Track 1 pillars.")

    # APPENDIX A4 (Slide 9 / 18)
    a4 = prs.slides.add_slide(blank)
    set_bg(a4, BG_LIGHT)
    add_chrome(a4, "APPENDIX A4 : ITT WORKFLOW & SPEND STOP", "Intention-to-Treat Cap-vs-Holdout & Late-Arrival Handling", "Appendix A4 (Slide 9 / 18)", True)
    steps_a4 = [
        ("STEP 1 : PASS ISSUANCE", "Filter 48 -> 40 eligible (80/20: 32 nudged, 8 control). Issue Rs 25 passes ('valid while shop offer budget lasts') to ALL 32 nudged buyers (Rs 800 face)."),
        ("STEP 2 : >10% & 80% ALERTS", "At approval, P(Binomial(32, 0.375) > 12) = 42.2% (>10% alert -> suggests Rs 400 cap). Live at 80% cap (10/12 paid), prompt merchant 1-tap top-up."),
        ("STEP 3 : INDEPENDENT DRAW", "48h simulator draws customer return visits independently of max_budget first, then applies atomic BEGIN IMMEDIATE stop at Rs 300 (12 paid)."),
        ("STEP 4 : ITT & DECLINE LOG", "After cap, normal UPI clears, cashback withholds, refused_redemptions + audit_log increment (brand risk tracked). ITT (32 vs 8) stays unchanged."),
    ]
    for i, (sh, sb) in enumerate(steps_a4):
        sx = 0.55 + i * 3.10
        add_box(a4, sx, 1.42, 2.92, 3.22, WHITE, border_rgb=NAVY, border_pt=1.2)
        add_text(a4, sx + 0.14, 1.56, 2.64, 0.36, sh, size=11.0, bold=True, color=NAVY)
        add_text(a4, sx + 0.14, 1.98, 2.64, 2.52, sb, size=10.8, color=INK)
    add_box(a4, 0.55, 4.82, 12.23, 2.04, AMB_SOFT, border_rgb=AMBER, border_pt=1.2)
    add_text(a4, 0.76, 4.94, 11.80, 0.32, "DISCLOSED BRAND RISK & VERIFIED INVARIANCE TESTS", size=12.5, bold=True, color=AMBER)
    add_bullet_list(
        a4, 0.72, 5.28, 11.85, 1.50,
        [
            ("Post-Cap Customer Experience & Brand Risk", "Pass terms state 'valid while shop offer budget lasts'; declining a 13th regular still risks counter disappointment, so overrun alerts, 80%-cap top-up prompts, and logged declines are mandatory."),
            ("Automated Concurrency & ITT Tests", "25-thread simultaneous scan test verifies budget_spent == Rs 300 (12 paid, 13 refused); multi-seed test proves ITT holdout comparison is identical whether cap is Rs 800 or Rs 100."),
        ],
        font_pt=11.2, space_pt=6
    )
    set_notes(a4, "Appendix A4 details pass validity terms, late-arrival brand risk, and automated concurrency plus ITT invariance tests.")

    # APPENDIX A5 (Slide 10 / 18)
    a5 = prs.slides.add_slide(blank)
    set_bg(a5, BG_LIGHT)
    add_chrome(a5, "APPENDIX A5 : SYSTEM ARCHITECTURE", "Component Status: Built vs Pending vs Production Plan", "Appendix A5 (Slide 10 / 18)", True)
    layers_a5 = [
        ("TOUCHPOINTS", "[BUILT] Merchant & Technical View UI", "[BUILT] Mic + Fallback Chips + Repeat/0.78x", "[PLANNED] Soundbox 4G Push & WhatsApp Cloud"),
        ("VOICE & LLM", "[BUILT] Browser Speech + SHA-256 Cache", "[PENDING LIVE] Keyless Sarvam (30% Cap)", "[BUILT/OPT] Keyless Local Q&A Guard"),
        ("CORE ENGINE", "[BUILT] FastAPI + TTL/LRU/Byte Sessions", "[BUILT] 6-Rule Guardrail Lab + Cap Editor", "[BUILT] ITT + Bayesian + Sensitivity Engine"),
        ("DATA & OPS", "[BUILT] Atomic SQLite Cap & Stop-Loss", "[BUILT] Preflight + Unit-Econ Generator", "[PLANNED] PostgreSQL & Redis Cluster"),
    ]
    for idx, (lname, b1, b2, b3) in enumerate(layers_a5):
        ly = 1.42 + idx * 1.36
        add_box(a5, 0.55, ly, 12.23, 1.20, WHITE, border_rgb=BORDER)
        add_box(a5, 0.55, ly, 2.20, 1.20, NAVY_DARK, rounded=False)
        add_text(a5, 0.68, ly + 0.38, 1.95, 0.45, lname, size=11.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        for j, btxt in enumerate([b1, b2, b3]):
            bx = 2.92 + j * 3.25
            is_built = btxt.startswith("[BUILT")
            is_pend = btxt.startswith("[PENDING")
            bg_c = EMER_SOFT if is_built else (AMB_SOFT if is_pend else SLATE_BG)
            bd_c = EMERALD if is_built else (AMBER if is_pend else BORDER)
            add_box(a5, bx, ly + 0.14, 3.10, 0.92, bg_c, border_rgb=bd_c, border_pt=1.1)
            add_text(a5, bx + 0.12, ly + 0.24, 2.86, 0.72, btxt, size=11.2, bold=True, color=INK)
    set_notes(a5, "Appendix A5 maps built Stage 1 and Stage 2 modules, pending live Saturday sponsor keys, and planned production hooks.")

    # APPENDIX A6 (Slide 11 / 18)
    a6 = prs.slides.add_slide(blank)
    set_bg(a6, BG_LIGHT)
    add_chrome(a6, "APPENDIX A6 : POSITIONING (USP)", "Differentiation vs Manual Khata Apps and Generic Chatbots", "Appendix A6 (Slide 11 / 18)", True)
    tbl_s = a6.shapes.add_table(6, 4, Inches(0.55), Inches(1.42), Inches(12.23), Inches(5.30))
    tbl = tbl_s.table
    for i, w in enumerate([Inches(2.35), Inches(3.00), Inches(3.00), Inches(3.88)]):
        tbl.columns[i].width = w
    hdrs = ["Dimension", "Manual BI / Khata Apps", "Generic LLM Chatbots", "PaytmAI VyaparVoice"]
    rows_a6 = [
        ["Merchant UX", "Charts & manual ledger entry", "Requires typing business prompts", "8:30 PM plain voice + Repeat/0.78x + 1-tap"],
        ["Customer Graph", "Incomplete manual phonebooks", "No link to payment history", "Paytm QR repeat-payer IDs (synthetic in MVP)"],
        ["Financial Safety", "Manual coupon rules", "LLM can invent unsafe offers", "6-rule guardrail lock, Cap Editor & Stop-Loss"],
        ["Privacy (DPDP)", "Raw phone lists exposed", "Prompt data leakage risk", "SQL opt-in filter; zero PII sent to voice/LLM"],
        ["Causal Proof", "Vanity message counts", "No comparison group", "80/20 ITT + Bayesian posterior + 50-shop pilot"],
    ]
    for c_i, h in enumerate(hdrs):
        cell = tbl.cell(0, c_i)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY_DARK if c_i < 3 else CYAN_DARK
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        p.text = assert_clean_text(h)
        p.font.name = FONT_MAIN
        p.font.size = Pt(12.0)
        p.font.bold = True
        p.font.color.rgb = WHITE
    for r_i, r_data in enumerate(rows_a6, start=1):
        for c_i, val in enumerate(r_data):
            cell = tbl.cell(r_i, c_i)
            cell.fill.solid()
            cell.fill.fore_color.rgb = EMER_SOFT if c_i == 3 else (WHITE if r_i % 2 == 1 else SLATE_BG)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.text = assert_clean_text(val)
            p.font.name = FONT_MAIN
            p.font.size = Pt(11.2)
            p.font.bold = c_i in (0, 3)
            p.font.color.rgb = INK
    set_notes(a6, "Appendix A6 contrasts VyaparVoice against manual dashboards and unconstrained LLM chatbots.")

    # APPENDIX A7 (Slide 12 / 18)
    a7 = prs.slides.add_slide(blank)
    set_bg(a7, BG_LIGHT)
    add_chrome(a7, "APPENDIX A7 : POOLED POWER & TARGETS", "50-Shop Stratified Power (400C + 1,600T) & Pilot Targets", "Appendix A7 (Slide 12 / 18)", True)
    add_box(a7, 0.55, 1.42, 6.50, 2.72, WHITE, border_rgb=BORDER)
    add_text(a7, 0.75, 1.52, 6.10, 0.28, "POOLED 50-SHOP X 1-WAVE POWER (12.5% BASELINE, ALPHA=0.05)", size=11.5, bold=True, color=NAVY)
    add_bullet_list(
        a7, 0.72, 1.84, 6.15, 2.22,
        [
            ("Why Pool Stores", "50 shops x 1 wave (~400C + ~1,600T, stratified by shop); avoids claiming 1 shop runs 4 waves in 4w under a 14d cooldown (~0.78 waves/mo)."),
            ("Exact Power Curve", "+4 pp: 50.3% | +5 pp: 68.8% | +6 pp: 83.3% | +7 pp: 92.5% | +8 pp: 97.2% | +10 pp: 99.8%."),
            ("Break-Even Lift", "+7.8 pp pre-fee (20.3%) and +15.3 pp post-5%-fee (27.8% at 30% shop margin; >99.9% power)."),
        ],
        font_pt=11.0, space_pt=5
    )
    add_box(a7, 0.55, 4.28, 6.50, 2.57, WHITE, border_rgb=BORDER)
    add_text(a7, 0.75, 4.38, 6.10, 0.28, "ECOSYSTEM HYPOTHESES (UNVALIDATED PILOT TARGETS)", size=11.5, bold=True, color=EMERALD)
    add_bullet_list(
        a7, 0.72, 4.70, 6.15, 2.05,
        [
            ("Repeat & Waste Targets", "+14% to +22% repeat visits; -18% afternoon perishable loss (PILOT TARGETS)."),
            ("Device & QR Share", "-35% Soundbox churn; +12% to +15% store Paytm QR GMV lift (PILOT TARGETS)."),
            ("Credit Funnel", "Expand ~7% merchant loan pen. (ASSUMPTION) toward 15% (LONG-TERM TARGET)."),
        ],
        font_pt=11.0, space_pt=5
    )
    add_box(a7, 7.25, 1.42, 5.53, 3.55, WHITE, border_rgb=BORDER)
    a7.shapes.add_picture(c5, Inches(7.52), Inches(1.48), width=Inches(4.98))
    tiles_a7 = [
        ("83.3%", "Power at +6 pp Lift\n(400C + 1,600T)", EMERALD),
        ("+15.3 pp", "Post-Fee Break-Even\n(+7.8 pp Pre-Fee)", AMBER),
        ("Rs 5.00", "Cached AI Cost/Mo\n(Rs 28.80 Uncached)", NAVY),
    ]
    for i, (tv, tl, tc) in enumerate(tiles_a7):
        tx = 7.25 + i * 1.88
        add_box(a7, tx, 5.14, 1.76, 1.70, WHITE, border_rgb=tc, border_pt=1.3)
        add_text(a7, tx + 0.06, 5.30, 1.64, 0.46, tv, size=17.5, bold=True, color=tc, align=PP_ALIGN.CENTER)
        add_text(a7, tx + 0.06, 5.84, 1.64, 0.86, tl, size=9.6, bold=True, color=MUTED, align=PP_ALIGN.CENTER)
    set_notes(a7, "Appendix A7 presents the exact power curve for pooling 50 shops times one wave stratified by shop.")

    # APPENDIX A8 (Slide 13 / 18)
    a8 = prs.slides.add_slide(blank)
    set_bg(a8, BG_LIGHT)
    add_chrome(a8, "APPENDIX A8 : UNIT ECONOMICS & FEE MATH", "Q2 FY27 ARR Cases, Fee Reconciliation & TTS Cache Sensitivity", "Appendix A8 (Slide 13 / 18)", True)
    add_box(a8, 0.55, 1.42, 6.50, 5.42, WHITE, border_rgb=BORDER)
    add_bullet_list(
        a8, 0.72, 1.54, 6.14, 5.20,
        [
            ("Stream 1 (SaaS - Rs 79/mo)", "Add-on to Soundbox (~Rs 90-100/mo rental ASSUMPTION) x 8.20L shops (5% of 1.64 Cr Q2 FY27 base) = Rs 77.7 Cr ARR."),
            ("Stream 2 (Promo Fee - Rs 102/mo)", "5% fee on ALL redeemed bills incl. organic returners: 1 wave = Rs 130.50 (Rs 2,610 bills); steady state assumes ~0.78 waves/mo (Rs 2,040/mo bills) = Rs 102/mo (+Rs 100.4 Cr -> Rs 178.1 Cr ARR)."),
            ("Stream 3 (Credit - Rs 77/mo)", "Amortized 1.5%-2.5% partner origination uplift (~7% loan pen. ASSUMPTION) -> +Rs 75.8 Cr (Rs 253.9 Cr ARR)."),
            ("Cached COGS (Rs 48.00/mo)", "Rs 1.80 SLM + Rs 3.20 TTS (~88% phrase cache hit on Bulbul v3) + Rs 43 msg = 81.4% gross margin."),
            ("Uncached COGS (Rs 71.80/mo)", "Uncached Bulbul v3 (9k chars at Rs 30/10k = Rs 27 TTS) = 72.2% margin (39.2% cached margin on SaaS-only)."),
        ],
        font_pt=11.2, space_pt=9
    )
    add_box(a8, 7.25, 1.42, 5.53, 3.60, WHITE, border_rgb=BORDER)
    a8.shapes.add_picture(c3, Inches(7.52), Inches(1.48), width=Inches(4.98))
    add_box(a8, 7.25, 5.18, 5.53, 1.66, NAVY_DARK, border_rgb=CYAN)
    add_text(
        a8, 7.42, 5.30, 5.18, 1.42,
        "REPRODUCIBLE SCRIPT (scripts/unit_economics.py)\n1. Fee Base: 5% applies to gross QR-redeemed promo bills including organic returners.\n2. Cadence: Initial backlog wave yields Rs 130.50; steady replenishment (~0.78 waves/mo) yields Rs 102/mo.\n3. Q2 FY27 Base: 5% x 1.64 Cr = 8.20L shops.",
        size=10.6, bold=True, color=WHITE
    )
    set_notes(a8, "Appendix A8 summarizes results/unit_economics.md generated reproducibly from config/default.yaml and app/stats.py.")

    # APPENDIX A9 (Slide 14 / 18)
    a9 = prs.slides.add_slide(blank)
    set_bg(a9, NAVY_DARK)
    add_box(a9, 0, 0, 13.333, 0.10, AMBER, rounded=False)
    add_text(a9, 0.65, 0.42, 7.5, 0.32, "APPENDIX A9 : VALIDATION & GO / NO-GO DECISION RULE", size=11.0, bold=True, color=CYAN)
    add_text(a9, 0.65, 0.85, 12.0, 0.50, "From Verified Stage 2 MVP to Pooled 50-Shop Stratified Pilot", size=24, bold=True, color=WHITE)
    phases_a9 = [
        ("STAGE A : BUILT & VERIFIED", "Deterministic Core & Stage 2 UX", [
            "Bounded visitor SQLite sandboxes (204 KB/visitor; 55 MB peak RSS)",
            "Cap Editor (Rs 50-2,500), 2-step voice confirm, Stop-Loss & Guardrail Lab",
            "Bayesian profit posterior, 6x4 sensitivity matrix & 5-point teach-back",
        ]),
        ("STAGE B : POOLED 50-SHOP PILOT", "Stratified Single-Wave Validation", [
            "3-merchant counter teach-back across EN (author-checked), KN & HI (draft)",
            "50 shops x 1 wave (~400C + ~1,600T, stratified by shop; 83.3% power at +6 pp)",
            "Advance only if lower 95% bound of post-5%-fee net profit/nudged > 0 AND lift > break-even (+15.3 pp)",
        ]),
        ("STAGE C : PRODUCTION SCALE", "Paytm Ecosystem Integration", [
            "Bind Paytm QR dynamic coupon settlement & 4G Soundbox push",
            "Enforce live 80%-cap merchant top-up prompt to protect counter goodwill",
            "Scale across 8.20L target stores (Rs 77.7 Cr to Rs 253.9 Cr ARR)",
        ]),
    ]
    for i, (pt, ph, pitems) in enumerate(phases_a9):
        px = 0.65 + i * 4.10
        add_box(a9, px, 1.65, 3.85, 3.88, NAVY, border_rgb=CYAN, border_pt=1.2)
        add_text(a9, px + 0.18, 1.80, 3.50, 0.30, pt, size=10.2, bold=True, color=CYAN)
        add_text(a9, px + 0.18, 2.14, 3.50, 0.36, ph, size=14.0, bold=True, color=WHITE)
        add_bullet_list(a9, px + 0.14, 2.56, 3.56, 2.82, pitems, font_pt=10.8, space_pt=9, lead_color=CYAN, body_color=RGBColor(0xE2, 0xE8, 0xF0))
    add_box(a9, 0.65, 5.75, 12.05, 1.10, RGBColor(0x00, 0x11, 0x33), border_rgb=CYAN)
    add_text(a9, 0.90, 5.92, 9.5, 0.36, "PaytmAI VyaparVoice | Team SemiTechnical (Kevin Joseph)", size=15.5, bold=True, color=WHITE)
    add_text(a9, 0.90, 6.32, 11.5, 0.34, "Appendix A9 (Slide 14 / 18) | Go/No-Go: Lower 95% CI of Post-5%-Fee Net Profit/Nudged > Rs 0 AND Observed Lift > Break-Even Lift", size=10.8, color=CYAN)
    set_notes(a9, "Appendix A9 states our pooled 50-shop pilot design and strict profit confidence-bound go/no-go decision rule.")

    # APPENDIX A10 (Slide 15 / 18)
    a10 = prs.slides.add_slide(blank)
    set_bg(a10, BG_LIGHT)
    add_chrome(a10, "APPENDIX A10 : DUAL SCORECARD & TEACH-BACK", "Merchant Plain-Language View vs Technical Audit View", "Appendix A10 (Slide 15 / 18)", True)
    add_box(a10, 0.55, 1.40, 5.95, 3.65, WHITE, border_rgb=EMERALD, border_pt=1.3)
    add_text(a10, 0.75, 1.54, 5.55, 0.32, "MERCHANT VIEW (ZERO JARGON - LINT VERIFIED)", size=12.5, bold=True, color=EMERALD)
    add_bullet_list(
        a10, 0.72, 1.92, 5.60, 2.95,
        [
            ("Plain Verdict & Why?", "Promising single-shop check (+Rs 222 pre-fee, +Rs 91.50 post-fee); 1-tap Why? drawers explain SQL triggers in plain words."),
            ("Counter Ergonomics", "Repeat Last, 0.78x Slower mode, and 5 Merchant Question Chips documented in docs/TEACH_BACK.md."),
            ("Clear Spend Stop", "Your 300 rupee limit was reached after 12 paid redemptions (0 extra declined in base run)."),
            ("Honest Copy Status", "EN author-checked, native review pending; KN & HI carry visible [DRAFT - UNREVIEWED] badges."),
        ],
        font_pt=11.2, space_pt=7
    )
    add_box(a10, 6.75, 1.40, 6.03, 3.65, WHITE, border_rgb=NAVY, border_pt=1.3)
    add_text(a10, 6.95, 1.54, 5.63, 0.32, "TECHNICAL & AUDIT VIEW (ANALYST / JUDGE)", size=12.5, bold=True, color=NAVY)
    add_bullet_list(
        a10, 6.92, 1.92, 5.68, 2.95,
        [
            ("ITT & Wilson 95% CIs", "37.5% [22.9%, 54.8%] vs 12.5% [2.2%, 47.1%]; two-sided Fisher exact p = 0.24."),
            ("Bayesian Posterior", "Beta(1,1) uniform prior (10k draws): 88.8% P(lift>0), 67.0% P(pre-fee>0), 52.1% P(post-fee>0)."),
            ("Guardrail Red-Team Lab", "Interactive 'Type your own attack' tests 6 hard rules (injection, <=22% discount, DPDP, 14d, cap, holdout)."),
            ("Banned Merchant Terms", "25 regex rules block cohort, holdout, p-value, Fisher, Wilson, ROI, DPDP, RFM, ITT."),
        ],
        font_pt=11.2, space_pt=7
    )
    grades_a10 = [
        ("GRADE A : OBSERVED LEDGER", "SQLite store counts (240 payers, 1,613 scans, 48 silent -> 40 eligible; synthetic seed 20261025).", NAVY),
        ("GRADE B : SIMULATED ITT", "Randomized 32 vs 8 split, independent outcome draw, Fisher p & Wilson CIs (simulated).", AMBER),
        ("GRADE C : ASSUMPTIONS", "30% shop gross margin, 5% fee on redeemed bills, Beta(1,1) profit posterior & sensitivity table.", EMERALD),
    ]
    for i, (gh, gb, gc) in enumerate(grades_a10):
        gx = 0.55 + i * 4.14
        add_box(a10, gx, 5.25, 3.94, 1.60, WHITE, border_rgb=gc, border_pt=1.3)
        add_text(a10, gx + 0.16, 5.38, 3.62, 0.30, gh, size=11.2, bold=True, color=gc)
        add_text(a10, gx + 0.16, 5.72, 3.62, 1.00, gb, size=10.8, color=INK)
    set_notes(a10, "Appendix A10 summarizes our Merchant View ergonomics, 5-point teach-back protocol, and Technical Audit View.")

    # APPENDIX A11 (Slide 16 / 18)
    a11 = prs.slides.add_slide(blank)
    set_bg(a11, BG_LIGHT)
    add_chrome(a11, "APPENDIX A11 : CAP LEDGERS & DOSE-RESPONSE", "Merchant Spend Stop, Sarvam Credit Cap & Dose-Response Curve", "Appendix A11 (Slide 16 / 18)", True)
    add_box(a11, 0.55, 1.40, 6.50, 5.45, WHITE, border_rgb=BORDER)
    add_bullet_list(
        a11, 0.72, 1.54, 6.14, 5.18,
        [
            ("Merchant Cashback Stop", "Issues passes ('valid while budget lasts') to all 32 nudged buyers; SQLite lock stops spend at Rs 300 (12 paid) and logs declines."),
            ("Dose-Response Sweet Spot", "In simulation, Rs 25 peaks uncapped profit: offers of Rs 30 to 40 erode uncapped net profit via organic subsidy AND trigger 5 to 8 cap refusals under a Rs 300 stop."),
            ("Margin & Lift Break-Even", "At +25 pp simulated lift, shop breaks even at 17.2% margin pre-fee (24.7% post-fee); at 30% margin, break-even lift is +7.8 pp pre-fee (+15.3 pp post-fee)."),
            ("Sarvam 30% Credit Ledger", "Separate JSON ledger caps TTS <=15% and STT <=15% (30% total of 500 sponsor credits = 150 max) and halts fail-closed on HTTP 402/429/5xx."),
            ("Public Cloud Lock", "PUBLIC_DEPLOY=true (auto-enabled on Render) hard-blocks live paid calls while SHA-256 audio cache replays at 0 credits."),
        ],
        font_pt=11.6, space_pt=10
    )
    add_box(a11, 7.25, 1.40, 5.53, 3.68, WHITE, border_rgb=BORDER)
    a11.shapes.add_picture(c4, Inches(7.46), Inches(1.48), width=Inches(5.10))
    add_box(a11, 7.25, 5.24, 5.53, 1.61, NAVY_DARK, border_rgb=CYAN)
    add_text(
        a11, 7.42, 5.36, 5.18, 1.38,
        "DUAL CAP SUMMARY (SIMULATED & BUILT)\n1. Merchant Cap: Protects shopkeeper pocket (Rs 300 stop; >10% overrun warning & 80% top-up prompt).\n2. Sponsor API Cap: Protects hackathon credits (<=15% TTS, <=15% STT, fail-closed).",
        size=10.8, bold=True, color=WHITE
    )
    set_notes(a11, "Appendix A11 shows how our simulated dose-response curve interacts with the Rs 300 spend stop and contrasts the merchant rupee cap against our separate 30 percent Sarvam credit ledger.")

    # APPENDIX A12 (Slide 17 / 18 - NEW STAGE 2 CAP EDITOR & STOP-LOSS SLIDE)
    a12 = prs.slides.add_slide(blank)
    set_bg(a12, BG_LIGHT)
    add_chrome(a12, "APPENDIX A12 : CAP EDITOR & LIVE STOP-LOSS [BUILT]", "Bounded Cap Editor, Two-Step Voice Confirm & Live Interim Stop-Loss", "Appendix A12 (Slide 17 / 18)", True)
    add_box(a12, 0.55, 1.40, 6.50, 5.45, WHITE, border_rgb=BORDER)
    add_bullet_list(
        a12, 0.72, 1.54, 6.14, 5.18,
        [
            ("Bounded Cap Editor [BUILT]", "Enforces hard bounds [Rs 50, Rs 2,500] (>=1 pass of Rs 25) with 1-tap presets: Rs 150, Rs 250, Rs 300 (base), Rs 400 (<=10% overrun), Rs 500."),
            ("Two-Step Voice Confirmation [BUILT]", "Saying 'Set my limit to 400 rupees' stages CAP_CONFIRMATION_PENDING without mutating SQLite until merchant says 'Confirm limit' or taps Confirm."),
            ("Live Interim Wave Monitor [BUILT]", "Stepping live scans displays an explicit INTERIM DIRECTIONAL READOUT banner so shopkeepers never stop a wave early on noise."),
            ("Atomic 1-Tap Stop-Loss [BUILT]", "Triggering Stop-Loss sets PAUSED_STOP_LOSS and blocks pass redemptions inside BEGIN IMMEDIATE (BLOCKED_STOP_LOSS)."),
            ("Immutable Audit Trail [BUILT]", "Every preset click, rejected out-of-bounds value, voice request/confirm, and Stop-Loss toggle writes to audit_logs."),
        ],
        font_pt=12.4, space_pt=12
    )
    add_box(a12, 7.25, 1.40, 5.53, 3.60, WHITE, border_rgb=BORDER)
    a12.shapes.add_picture(c6, Inches(7.50), Inches(1.46), width=Inches(5.02))
    add_box(a12, 7.25, 5.16, 5.53, 1.69, NAVY_DARK, border_rgb=CYAN)
    add_text(
        a12, 7.42, 5.28, 5.18, 1.46,
        "CAP PRESETS VS BINOMIAL REFUSAL RISK (p = 37.5%, 32 PASSES)\nRs 150 (6 paid): 98.1% overrun | Rs 250 (10 paid): 70.4%\nRs 300 (12 paid): 42.2% (>10% alert) | Rs 400 (16 paid): 5.2% (SAFE)\nRs 500 (20 paid): 0.1% overrun.",
        size=10.4, bold=True, color=WHITE
    )
    set_notes(a12, "Appendix A12 documents our Stage 2 Task 9 Cap Editor, two-step voice cap confirmation, live interim readout banner, and atomic SQLite Stop-Loss guard.")

    # APPENDIX A13 (Slide 18 / 18 - NEW STAGE 2 BAYESIAN SCORECARD & SENSITIVITY SLIDE)
    a13 = prs.slides.add_slide(blank)
    set_bg(a13, BG_LIGHT)
    add_chrome(a13, "APPENDIX A13 : BAYESIAN SCORECARD & SENSITIVITY [BUILT]", "Beta(1,1) Profit Posterior, Evidence Meter & 6x4 Margin-Fee Matrix", "Appendix A13 (Slide 18 / 18)", True)
    add_box(a13, 0.55, 1.40, 6.50, 5.45, WHITE, border_rgb=BORDER)
    add_bullet_list(
        a13, 0.72, 1.54, 6.14, 5.18,
        [
            ("Plain-Language Verdict [BUILT]", "Merchant View translates single-wave output into plain rupee terms with zero statistical jargon, backed by Grade A/B/C badges."),
            ("Bayesian Profit Posterior [BUILT]", "Conjugate Beta(1,1) prior (10k draws, seed 20261025) without clamping negative draws: 88.8% P(lift>0), 67.0% P(pre-fee>0), 52.1% P(post-fee>0)."),
            ("Why Single Wave Cannot Auto-Scale", "Single-wave 95% CrI is [-Rs 19.97, +Rs 14.73] per nudged buyer (due to n=8 control uncertainty), failing A9 scale gate until 50 shops are pooled."),
            ("4-Stage Visual Evidence Meter [BUILT]", "Tracks progress from n=40 (1 shop) -> n=155 (+25pp power) -> n=600 (15 shops) -> n=2,000 (50-shop stratified pilot, 83.3% power at +6 pp)."),
            ("Kirana Margin Sensitivity [BUILT]", "At 5% fee, Rs 25 pass profits at >=25% margin (+Rs 4.50 at 25%, +Rs 91.50 at 30%) and loses at 15%-20% (-Rs 82.50 at 20%), guiding smaller kirana offers."),
        ],
        font_pt=12.2, space_pt=11
    )
    add_box(a13, 7.25, 1.40, 5.53, 3.65, WHITE, border_rgb=BORDER)
    a13.shapes.add_picture(c7, Inches(7.44), Inches(1.46), width=Inches(5.14))
    add_box(a13, 7.25, 5.20, 5.53, 1.65, EMER_SOFT, border_rgb=EMERALD, border_pt=1.2)
    add_text(
        a13, 7.42, 5.30, 5.18, 1.45,
        "HONEST SCORECARD TAKEAWAY (TRACED TO unit_economics.md)\nBase [30% Margin, 5% Fee] = +Rs 91.50 (+Rs 222 pre-fee).\nLow-margin grocery [15%-20% Margin] requires lowering pass to Rs 10-15 on Rs 250+ bills (4%-6% discount).",
        size=10.5, bold=True, color=INK
    )
    set_notes(a13, "Appendix A13 shows our built Stage 2 Task 11 Merchant Scorecard: the Beta(1,1) profit posterior, 4-stage Evidence Meter, and 6 by 4 Shop Gross Margin times Paytm Fee sensitivity matrix.")

    prs.save(PPTX_V3_OUT)
    shutil.copy2(PPTX_V3_OUT, os.path.join(OUT_DIRS[1], "PaytmAI_VyaparVoice_v3.pptx"))
    print("Saved Deck v3:", PPTX_V3_OUT)


if __name__ == "__main__":
    build_deck_v3()
