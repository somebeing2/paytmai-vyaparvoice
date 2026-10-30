"""
Deck v1 Generator for PaytmAI VyaparVoice.
Produces a 5-slide Lean Main Deck (>=18pt body text, <=5 bullets/slide)
followed by a 9-slide Corrected Appendix (Slides 6-14).
Strictly enforces: zero em dashes, zero en dashes, zero emojis, full traceability
to results/slide_numbers.md and AUDIT.md.
"""
import os
import re
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ASSET_DIR = "/home/user/scripts/v1_assets"
OUT_DIRS = [
    "/home/user/paytmai-vyaparvoice/deck/v1",
    "/home/user/deck/v1",
]
for d in [ASSET_DIR] + OUT_DIRS:
    os.makedirs(d, exist_ok=True)

PPTX_OUT = os.path.join(OUT_DIRS[0], "PaytmAI_VyaparVoice_v1.pptx")

# Palette
NAVY_DARK = RGBColor(0x00, 0x1B, 0x4D)
NAVY      = RGBColor(0x00, 0x29, 0x70)
CYAN      = RGBColor(0x00, 0xBA, 0xF2)
CYAN_DARK = RGBColor(0x02, 0x84, 0xC7)
CYAN_SOFT = RGBColor(0xE0, 0xF7, 0xFE)
BG_LIGHT  = RGBColor(0xF5, 0xF8, 0xFC)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
INK       = RGBColor(0x0F, 0x17, 0x2A)
MUTED     = RGBColor(0x33, 0x41, 0x55)
BORDER    = RGBColor(0xCB, 0xD5, 0xE1)
EMERALD   = RGBColor(0x05, 0x96, 0x69)
EMER_SOFT = RGBColor(0xEC, 0xFD, 0xF5)
AMBER     = RGBColor(0xB4, 0x53, 0x09)
AMB_SOFT  = RGBColor(0xFF, 0xFB, 0xEB)
SLATE_BG  = RGBColor(0xF1, 0xF5, 0xF9)

FONT_MAIN = "Liberation Sans"

EMOJI_RE = re.compile(
    "["
    "\U0001F300-\U0001FAFF"
    "\U00002700-\U000027BF"
    "\U00002600-\U000026FF"
    "]+"
)


def assert_clean_text(text: str) -> str:
    if "\u2014" in text or "\u2013" in text:
        raise ValueError(f"Banned dash found in text: {text!r}")
    if EMOJI_RE.search(text):
        raise ValueError(f"Banned emoji found in text: {text!r}")
    return text


def generate_v1_charts():
    plt.rcParams["font.sans-serif"] = ["DejaVu Sans", "Liberation Sans", "Arial"]
    plt.rcParams["axes.edgecolor"] = "#CBD5E1"
    plt.rcParams["axes.linewidth"] = 0.9

    # Chart 1: Actual 30-day hourly scans from synthetic SQLite ledger (seed 20261025)
    import sys
    sys.path.insert(0, "/home/user/paytmai-vyaparvoice")
    from app.database import init_and_seed_db
    from app.engine import get_dashboard_overview
    init_and_seed_db(force_reset=False)
    ov = get_dashboard_overview()
    hrs = ov["hourly_series"]
    labels = [h["label"] for h in hrs]
    counts = [h["txns"] for h in hrs]
    cols = ["#D97706" if h["is_dead_hour"] else "#002970" for h in hrs]

    fig, ax = plt.subplots(figsize=(5.1, 3.6), dpi=260)
    fig.patch.set_facecolor("#FFFFFF")
    ax.set_facecolor("#FFFFFF")
    ax.bar(labels, counts, color=cols, width=0.65, zorder=3)
    ax.grid(axis="y", linestyle="--", alpha=0.45, color="#CBD5E1", zorder=0)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.set_ylabel("30-Day UPI Scans (Count)", fontsize=9.5, fontweight="bold", color="#1E293B")
    ax.set_title(assert_clean_text("Synthetic Store 30d Hourly Scans (n = 917)"), fontsize=10.5, fontweight="bold", color="#002970", pad=10)
    ax.tick_params(axis="x", labelsize=7.5, colors="#334155", rotation=30)
    ax.tick_params(axis="y", labelsize=8.5, colors="#334155")
    ax.set_ylim(0, 122)
    ax.annotate(
        assert_clean_text("3-5 PM Dead-Hour\n(30 scans vs 101 peak = -70%)"),
        xy=(8.0, 31),
        xytext=(4.2, 102),
        arrowprops=dict(arrowstyle="->", color="#B45309", lw=1.5),
        fontsize=8.2,
        fontweight="bold",
        color="#B45309",
        bbox=dict(boxstyle="round,pad=0.3", fc="#FFFBEB", ec="#D97706", lw=1),
    )
    plt.tight_layout()
    c1 = os.path.join(ASSET_DIR, "c1_hourly_synthetic.png")
    plt.savefig(c1, dpi=260)
    plt.close(fig)

    # Chart 2: Honest Holdout Comparison with Wilson 95% CIs (SIMULATED)
    fig, ax = plt.subplots(figsize=(5.1, 3.6), dpi=260)
    fig.patch.set_facecolor("#FFFFFF")
    ax.set_facecolor("#FFFFFF")
    arms = [assert_clean_text("Control Arm\n(1/8 = 12.5%)"), assert_clean_text("Nudged Arm\n(12/32 = 37.5%)")]
    means = [12.5, 37.5]
    # Wilson 95% CIs: Control [2.24, 47.09], Treatment [22.93, 54.75]
    lower_err = [12.5 - 2.24, 37.5 - 22.93]
    upper_err = [47.09 - 12.5, 54.75 - 37.5]
    bars = ax.bar(arms, means, color=["#64748B", "#059669"], width=0.45, zorder=3)
    ax.errorbar(arms, means, yerr=[lower_err, upper_err], fmt="none", ecolor="#0F172A", elinewidth=1.8, capsize=7, zorder=4, label="Wilson 95% CI")
    ax.grid(axis="y", linestyle="--", alpha=0.45, color="#CBD5E1", zorder=0)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.set_ylabel("48h Return Rate (%)", fontsize=9.5, fontweight="bold", color="#1E293B")
    ax.set_title(assert_clean_text("SIMULATED Holdout + 95% Wilson CIs (p = 0.24)"), fontsize=10.2, fontweight="bold", color="#002970", pad=10)
    ax.set_ylim(0, 68)
    ax.tick_params(axis="both", labelsize=9.0, colors="#1E293B")
    ax.text(0, 49.5, "95% CI:\n2.2%-47.1%", ha="center", fontsize=8.0, color="#334155", fontweight="bold")
    ax.text(1, 57.0, "95% CI:\n22.9%-54.8%", ha="center", fontsize=8.0, color="#065F46", fontweight="bold")
    ax.legend(loc="upper left", fontsize=8.2, frameon=True, facecolor="#F8FAFC")
    plt.tight_layout()
    c2 = os.path.join(ASSET_DIR, "c2_honest_holdout.png")
    plt.savefig(c2, dpi=260)
    plt.close(fig)

    # Chart 3: Three Revenue Cases at 5% Soundbox Adoption (7.85L Merchants, MODELLED)
    fig, ax = plt.subplots(figsize=(5.1, 3.6), dpi=260)
    fig.patch.set_facecolor("#FFFFFF")
    ax.set_facecolor("#FFFFFF")
    cases = [
        assert_clean_text("Case C:\nSaaS Only"),
        assert_clean_text("Case B:\nSaaS + Promo"),
        assert_clean_text("Case A:\nFull Base"),
    ]
    saas = np.array([74.4, 74.4, 74.4])
    promo = np.array([0.0, 96.1, 96.1])
    credit = np.array([0.0, 0.0, 72.5])
    x = np.arange(3)
    w = 0.48
    ax.bar(x, saas, width=w, label="Rs 79/mo SaaS", color="#002970", zorder=3)
    ax.bar(x, promo, width=w, bottom=saas, label="5% Promo Fee (Rs 102/mo)", color="#0284C7", zorder=3)
    ax.bar(x, credit, width=w, bottom=saas + promo, label="Credit Referral (Rs 77/mo)", color="#059669", zorder=3)
    tots = saas + promo + credit
    for xi, tot in zip(x, tots):
        ax.text(xi, tot + 6, f"Rs {tot:.1f} Cr", ha="center", va="bottom", fontsize=9.5, fontweight="bold", color="#0F172A")
    ax.set_xticks(x)
    ax.set_xticklabels(cases, fontsize=8.8, fontweight="bold", color="#1E293B")
    ax.set_ylabel("Modelled ARR (Rs Crore)", fontsize=9.5, fontweight="bold", color="#1E293B")
    ax.set_title(assert_clean_text("Modelled ARR at 5% Base (7.85L Shops)"), fontsize=10.3, fontweight="bold", color="#002970", pad=10)
    ax.set_ylim(0, 290)
    ax.grid(axis="y", linestyle="--", alpha=0.45, color="#CBD5E1", zorder=0)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.legend(loc="upper left", fontsize=7.8, frameon=True, facecolor="#F8FAFC")
    plt.tight_layout()
    c3 = os.path.join(ASSET_DIR, "c3_three_revenue_cases.png")
    plt.savefig(c3, dpi=260)
    plt.close(fig)

    return c1, c2, c3


def set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_box(slide, left, top, width, height, fill_rgb, border_rgb=None, border_pt=1.0, rounded=True):
    st = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(st, Inches(left), Inches(top), Inches(width), Inches(height))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill_rgb
    if border_rgb:
        shp.line.color.rgb = border_rgb
        shp.line.width = Pt(border_pt)
    else:
        shp.line.fill.background()
    if rounded and len(shp.adjustments) > 0:
        shp.adjustments[0] = 0.03
    return shp


def add_text(slide, left, top, width, height, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, italic=False):
    assert_clean_text(text)
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(3)
        r = p.add_run()
        r.text = line
        r.font.name = FONT_MAIN
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
    return tb


def add_bullet_list(slide, left, top, width, height, bullets, font_pt=18.0, space_pt=10):
    if len(bullets) > 5:
        raise ValueError(f"Exceeded 5 bullets limit: got {len(bullets)}")
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    for idx, (lead, body) in enumerate(bullets):
        assert_clean_text(lead)
        assert_clean_text(body)
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.space_after = Pt(space_pt)
        r1 = p.add_run()
        r1.text = f"* {lead}: "
        r1.font.name = FONT_MAIN
        r1.font.size = Pt(font_pt)
        r1.font.bold = True
        r1.font.color.rgb = NAVY

        r2 = p.add_run()
        r2.text = body
        r2.font.name = FONT_MAIN
        r2.font.size = Pt(font_pt)
        r2.font.bold = False
        r2.font.color.rgb = INK
    return tb


def add_chrome(slide, tag_text: str, title: str, page_label: str, is_appendix: bool = False):
    assert_clean_text(tag_text)
    assert_clean_text(title)
    assert_clean_text(page_label)
    add_box(slide, 0, 0, 13.333, 0.08, AMBER if is_appendix else CYAN, rounded=False)

    pill_bg = AMB_SOFT if is_appendix else CYAN_SOFT
    pill_fg = AMBER if is_appendix else NAVY
    pill_w = max(2.6, len(tag_text) * 0.095 + 0.5)
    add_box(slide, 0.55, 0.24, pill_w, 0.32, pill_bg, border_rgb=pill_fg, border_pt=1.0)
    add_text(slide, 0.60, 0.27, pill_w - 0.1, 0.26, tag_text, size=9.5, bold=True, color=pill_fg, align=PP_ALIGN.CENTER)

    add_text(
        slide, 7.6, 0.26, 5.18, 0.28,
        "PaytmAI VyaparVoice | Team SemiTechnical (Kevin Joseph)",
        size=9.5, bold=True, color=MUTED, align=PP_ALIGN.RIGHT
    )
    add_text(slide, 0.55, 0.66, 12.23, 0.52, title, size=23.5, bold=True, color=INK)

    add_box(slide, 0.55, 7.08, 12.23, 0.015, BORDER, rounded=False)
    add_text(
        slide, 0.55, 7.13, 9.2, 0.26,
        "Paytm Build for India AI Hackathon (Bengaluru Track 1) | Figures traced to results/slide_numbers.md",
        size=8.8, color=MUTED
    )
    add_text(slide, 10.6, 7.13, 2.18, 0.26, page_label, size=8.8, bold=True, color=NAVY, align=PP_ALIGN.RIGHT)


def set_notes(slide, notes_text: str):
    assert_clean_text(notes_text)
    slide.notes_slide.notes_text_frame.text = notes_text


def build_deck_v1():
    c1, c2, c3 = generate_v1_charts()
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # =========================================================================
    # MAIN SLIDE 1 / 5: PROBLEM (WITH PROJECT & TEAM HEADER)
    # =========================================================================
    m1 = prs.slides.add_slide(blank)
    set_bg(m1, BG_LIGHT)
    add_chrome(
        m1,
        "MAIN 1 / 5 : PROBLEM & OPPORTUNITY",
        "1.57 Cr Storefronts Accept QR Payments; Retention Stays Manual",
        "Main 1 / 5 (Slide 1 / 14)",
    )

    add_box(m1, 0.55, 1.30, 7.25, 5.55, WHITE, border_rgb=BORDER, border_pt=1.0)
    add_bullet_list(
        m1, 0.72, 1.48, 6.90, 5.20,
        [
            ("Verified Paytm Scale", "1.57 Cr merchant device subscribers and Rs 7.1L Cr quarterly GMV (Paytm Q1 FY27, Jul 2026)."),
            ("Silent Repeat Churn", "Small shops lose an estimated 25% to 35% of repeat buyers yearly without alerts (Assumption)."),
            ("No Time for Dashboards", "Counter operators work 12+ hr days; SMB BI apps see <8% weekly active usage (Assumption)."),
            ("Dead-Hour & Prep Gaps", "Afternoon scans (3 to 5 PM) drop 70% vs peak in our synthetic store, hurting daily margins."),
        ],
        font_pt=18.2, space_pt=14
    )

    add_box(m1, 8.00, 1.30, 4.78, 3.75, WHITE, border_rgb=BORDER, border_pt=1.0)
    m1.shapes.add_picture(c1, Inches(8.08), Inches(1.40), width=Inches(4.62))

    add_box(m1, 8.00, 5.20, 4.78, 1.65, NAVY_DARK, border_rgb=CYAN, border_pt=1.2)
    add_text(
        m1, 8.18, 5.32, 4.42, 1.42,
        "PaytmAI VyaparVoice (Track 1)\nProactive closing-time voice copilot for Paytm Soundbox & Business App.\nTeam SemiTechnical | Kevin Joseph",
        size=13.5, bold=True, color=WHITE
    )

    set_notes(
        m1,
        "Good morning. I am Kevin Joseph from Team SemiTechnical, presenting PaytmAI VyaparVoice for Track 1, Merchant Growth AI. "
        "Per One97 Communications Q1 FY27 disclosures from July 2026, Paytm has 1.57 crore merchant device subscriptions processing Rs 7.1 lakh crore in quarterly GMV. "
        "However, small shopkeepers still have no automated retention system: we estimate 25 to 35 percent annual repeat-buyer churn, and dashboard tools fail because counter operators do not type prompts during a 12-hour shift. "
        "In our synthetic Bengaluru darshini dataset, afternoon scans between 3 and 5 PM also drop 70 percent below morning peak."
    )

    # =========================================================================
    # MAIN SLIDE 2 / 5: SOLUTION & LIVE DEMO MOMENT
    # =========================================================================
    m2 = prs.slides.add_slide(blank)
    set_bg(m2, BG_LIGHT)
    add_chrome(
        m2,
        "MAIN 2 / 5 : SOLUTION & LIVE DEMO",
        "One Closing-Time Voice Action, Governed by Deterministic Rules",
        "Main 2 / 5 (Slide 2 / 14)",
    )

    add_box(m2, 0.55, 1.30, 12.23, 5.55, WHITE, border_rgb=BORDER, border_pt=1.0)
    add_bullet_list(
        m2, 0.78, 1.50, 11.75, 5.15,
        [
            ("8:30 PM Voice Pitch", "Speaks one concrete action in Kannada, Hindi, or English (draft strings; native review pending)."),
            ("Synthetic Store Demo", "Sri Krishna Darshini (240 payers, 1,613 UPI scans); detects 48 regulars silent 16+ days (median 19d)."),
            ("5-Rule Policy Engine", "Excludes 5 unconsented (DPDP) and 3 on 14d cooldown; splits 40 eligible into 32 nudged / 8 control."),
            ("Money Guardrails", "Rs 25 cashback on Rs 150+ scan (16.7% <= 22% floor), Rs 300 cap; no LLM ever writes numbers or caps."),
            ("Ops & Care Actions", "Same 1-tap flow triggers simulated weekend dairy restock POs and queue-delay customer recovery."),
        ],
        font_pt=18.2, space_pt=13
    )

    set_notes(
        m2,
        "Instead of asking merchants to read charts, VyaparVoice speaks one actionable recommendation at 8:30 PM closing time. "
        "In our live MVP demo, which runs on a synthetic dataset of 240 payers and 1,613 UPI scans for Sri Krishna Darshini in Indiranagar, the SQL engine detects 48 regulars silent for at least 16 days, with a median gap of 19 days. "
        "Before pitching, a 5-rule deterministic Python engine filters out 5 unconsented payers and 3 payers inside a 14-day anti-spam cooldown, leaving 40 eligible customers split 32 treatment and 8 silent control. "
        "The merchant approves by button or voice, locking a Rs 25 cashback on bills of Rs 150 or more within a Rs 300 budget cap."
    )

    # =========================================================================
    # MAIN SLIDE 3 / 5: HOW IT WORKS (BUILT VS PLANNED)
    # =========================================================================
    m3 = prs.slides.add_slide(blank)
    set_bg(m3, BG_LIGHT)
    add_chrome(
        m3,
        "MAIN 3 / 5 : ARCHITECTURE (BUILT VS PLANNED)",
        "Deterministic Core Built Today; Ecosystem Integrations Scoped",
        "Main 3 / 5 (Slide 3 / 14)",
    )

    add_box(m3, 0.55, 1.30, 12.23, 5.55, WHITE, border_rgb=BORDER, border_pt=1.0)
    add_bullet_list(
        m3, 0.78, 1.50, 11.75, 5.15,
        [
            ("[BUILT IN MVP] FastAPI + SQLite Core", "Live SQL cohort queries, 5-rule policy checks, 80/20 holdout math, audit log, 21/21 QA tests."),
            ("[BUILT IN MVP] Companion PWA + Audio", "Paytm Soundbox simulator UI using browser speech synthesis + 1-tap button fallback (offline-safe)."),
            ("[PENDING VERIFICATION] Sponsor Voice & Q&A", "Keyless adapters for Sarvam TTS/ASR (Saturday window) and free-tier Groq grounded English Q&A."),
            ("[PRODUCTION PLAN] Paytm Rail Hooks", "Paytm QR dynamic coupon ledger, Soundbox 4G push audio, WhatsApp Cloud POs, Postgres/Redis."),
        ],
        font_pt=18.2, space_pt=15
    )

    set_notes(
        m3,
        "We strictly separate what is built and running today from our production integration plan. "
        "Built today in our FastAPI and SQLite MVP are live SQL cohort segmentation, the 5-rule deterministic policy engine, the 80/20 holdout calculator, an immutable audit log, and a browser-based Soundbox companion UI with 21 passing automated checks. "
        "Sarvam AI voice synthesis and speech recognition are wired via keyless adapters and marked pending until verified during tomorrow's sponsor build window. "
        "Integration into Paytm's live QR coupon settlement rail, 4G Soundbox push firmware, WhatsApp Business Cloud, and Redis/Postgres are scoped for the production phase."
    )

    # =========================================================================
    # MAIN SLIDE 4 / 5: PROOF WITH HONEST STATISTICS (SIMULATED)
    # =========================================================================
    m4 = prs.slides.add_slide(blank)
    set_bg(m4, BG_LIGHT)
    add_chrome(
        m4,
        "MAIN 4 / 5 : SIMULATED PROOF & HONEST STATS",
        "Simulated Holdout Lift & Profit Math (To Validate in Pilot)",
        "Main 4 / 5 (Slide 4 / 14)",
    )

    add_box(m4, 0.55, 1.30, 7.25, 5.55, WHITE, border_rgb=BORDER, border_pt=1.0)
    add_bullet_list(
        m4, 0.72, 1.44, 6.92, 5.30,
        [
            ("Simulated 48h Return", "12/32 nudged (37.5%) vs 1/8 control (12.5%); Fisher exact p = 0.24 (not significant)."),
            ("Pilot Sample Target", "80% power at alpha = 0.05 needs ~46 payers/arm (~4 weekly store waves)."),
            ("Sales Multiple (5.8x)", "8 net visits x Rs 217.50 assumed bill = Rs 1,740 incremental sales on Rs 300 spend."),
            ("Organic Subsidy", "4 of 12 redeemers (12.5% baseline) would visit anyway, absorbing Rs 100 cashback."),
            ("Profit ROI (30% Margin)", "Rs 522 gross profit (1.74x); +Rs 222 net pre-fee (+74% ROI), +Rs 91.50 post 5% fee."),
        ],
        font_pt=18.0, space_pt=10
    )

    add_box(m4, 8.00, 1.30, 4.78, 3.85, WHITE, border_rgb=BORDER, border_pt=1.0)
    m4.shapes.add_picture(c2, Inches(8.08), Inches(1.40), width=Inches(4.62))

    add_box(m4, 8.00, 5.30, 4.78, 1.55, AMB_SOFT, border_rgb=AMBER, border_pt=1.2)
    add_text(
        m4, 8.16, 5.42, 4.46, 1.30,
        "EVIDENCE DISCLOSURE (SIMULATED)\nRedemptions come from our deterministic 48h simulator. Break-even store gross margin is 17.2% pre-fee (24.7% after 5% Paytm promo fee).",
        size=12.5, bold=True, color=INK
    )

    set_notes(
        m4,
        "We report our holdout simulation with full statistical and accounting transparency. "
        "In our simulated 48-hour run, 12 of 32 nudged regulars return versus 1 of 8 in control. With N equals 40, Fisher's exact two-sided p-value is 0.24, and the 95 percent Wilson intervals overlap, so this single wave is directional only; reaching 80 percent power requires roughly 46 customers per arm across 4 weekly waves. "
        "On accounting: 8 net incremental buyers spending an assumed Rs 217.50 generate Rs 1,740 in incremental sales, so 5.8x is gross sales per rupee of cashback, while 4 organic returners absorb Rs 100 of cashback. "
        "At a 30 percent gross margin assumption, incremental gross profit is Rs 522, leaving Rs 222 net profit before fees, or Rs 91.50 net profit after a 5 percent Paytm fee on all redeemed bills."
    )

    # =========================================================================
    # MAIN SLIDE 5 / 5: BUSINESS MODEL & PILOT PLAN
    # =========================================================================
    m5 = prs.slides.add_slide(blank)
    set_bg(m5, BG_LIGHT)
    add_chrome(
        m5,
        "MAIN 5 / 5 : BUSINESS MODEL & PILOT PLAN",
        "Three Modelled Revenue Cases (5% Base) & 45-Day Pilot Plan",
        "Main 5 / 5 (Slide 5 / 14)",
    )

    add_box(m5, 0.55, 1.30, 7.25, 5.55, WHITE, border_rgb=BORDER, border_pt=1.0)
    add_bullet_list(
        m5, 0.72, 1.44, 6.92, 5.30,
        [
            ("Case C (SaaS Only)", "Rs 79/mo Soundbox add-on x 7.85L shops (5% base) = Rs 74.4 Cr ARR (Modelled)."),
            ("Case B (Excl. Credit)", "Adding 5% fee on redeemed bills (+Rs 102/mo) = Rs 170.5 Cr ARR (Modelled)."),
            ("Case A (Full Base)", "Adding partner loan referral (+Rs 77/mo; Paytm pen. ~7%, Q4 FY26) = Rs 243.0 Cr ARR."),
            ("Consistent Unit COGS", "Rs 48/mo cached (Rs 5 AI + Rs 43 msg; 81.4% margin) or Rs 71.80 uncached TTS (72.2%)."),
            ("45-Day Bengaluru Pilot", "50 kirana/F&B shops, >=50 payers/arm pooled holdout, native Kannada/Hindi sign-off."),
        ],
        font_pt=18.0, space_pt=10
    )

    add_box(m5, 8.00, 1.30, 4.78, 3.85, WHITE, border_rgb=BORDER, border_pt=1.0)
    m5.shapes.add_picture(c3, Inches(8.08), Inches(1.40), width=Inches(4.62))

    add_box(m5, 8.00, 5.30, 4.78, 1.55, EMER_SOFT, border_rgb=EMERALD, border_pt=1.2)
    add_text(
        m5, 8.16, 5.42, 4.46, 1.30,
        "TTS CACHE ASSUMPTION\nBulbul v3 lists Rs 30 / 10k chars (~Rs 27/mo uncached at 300 chars/day). Achieving Rs 3.20/mo TTS requires ~88% template caching (or v2 pricing).",
        size=12.2, bold=True, color=INK
    )

    set_notes(
        m5,
        "Because Paytm already bills 1.57 crore device merchants monthly, distribution rides an existing rail. "
        "Rather than relying on a single optimistic projection, we model three cases at 5 percent adoption, or 7.85 lakh shops: a conservative SaaS-only case at Rs 79 per month yielding Rs 74.4 crore ARR; a core product case adding a 5 percent fee on redeemed promo bills yielding Rs 170.5 crore ARR; and a full ecosystem case adding partner lending origination yielding Rs 243.0 crore ARR. "
        "Unit COGS is Rs 48 per month assuming 88 percent TTS phrase caching for an 81.4 percent gross margin, or Rs 71.80 per month uncached for a 72.2 percent margin. "
        "Next is a 45-day, 50-store Bengaluru pilot powered to validate true causal lift."
    )

    # =========================================================================
    # APPENDIX SLIDES (SLIDES 6 TO 14 - CORRECTED ORIGINAL 9 SLIDES)
    # =========================================================================

    # A1 (Slide 6 / 14): Corrected Title & Overview Card
    a1 = prs.slides.add_slide(blank)
    set_bg(a1, NAVY_DARK)
    add_box(a1, 0, 0, 13.333, 0.10, AMBER, rounded=False)
    add_box(a1, 0.65, 0.45, 7.60, 0.42, NAVY, border_rgb=CYAN, border_pt=1.0)
    add_text(a1, 0.80, 0.52, 7.30, 0.28, "APPENDIX A1 (CORRECTED ORIGINAL SLIDE 1) | TRACK 1: MERCHANT GROWTH AI", size=10.0, bold=True, color=CYAN)
    add_text(a1, 0.65, 1.25, 7.6, 0.80, "PaytmAI VyaparVoice", size=40, bold=True, color=WHITE)
    add_text(a1, 0.65, 2.10, 7.6, 0.70, "Proactive Voice Growth, Operations & Retention Copilot for Paytm Merchants", size=17.5, bold=True, color=CYAN)
    add_text(
        a1, 0.65, 2.95, 7.5, 0.90,
        "Designed for Paytm's 1.57 Cr merchant device base (Q1 FY27). Combines deterministic SQL cohort rules and budget guardrails with closing-time voice briefings (synthetic MVP built; live pilot pending).",
        size=13.0, color=RGBColor(0xCB, 0xD5, 0xE1)
    )
    cards_a1 = [
        ("01 : Voice + Button Flow", "English, Kannada & Hindi draft scripts; 1-tap button fallback"),
        ("02 : 5-Rule Policy Engine", "DPDP opt-in, 14d cooldown, <=22% floor & Rs 300 cap"),
        ("03 : Simulated Holdout", "80/20 control split; 5.8x sales/promo (1.74x gross profit)"),
    ]
    for i, (ch, cb) in enumerate(cards_a1):
        bx = 0.65 + i * 2.55
        add_box(a1, bx, 4.15, 2.40, 1.15, NAVY, border_rgb=CYAN, border_pt=1.0)
        add_text(a1, bx + 0.12, 4.25, 2.16, 0.30, ch, size=10.8, bold=True, color=WHITE)
        add_text(a1, bx + 0.12, 4.58, 2.16, 0.60, cb, size=9.8, color=RGBColor(0xBA, 0xE6, 0xFD))
    add_box(a1, 0.65, 5.65, 12.0, 1.10, RGBColor(0x00, 0x12, 0x35), border_rgb=CYAN, border_pt=1.0)
    add_text(a1, 0.85, 5.80, 6.0, 0.35, "TEAM: SemiTechnical | Builder: Kevin Joseph (Bengaluru)", size=14, bold=True, color=WHITE)
    add_text(a1, 0.85, 6.20, 11.4, 0.35, "Appendix Status: Corrected from AUDIT.md (all synthetic, simulated, and pending items explicitly labelled).", size=11.5, color=CYAN)
    set_notes(a1, "Appendix A1 preserves the original title slide layout with all unverified claims removed and 5.8x clarified as simulated incremental sales per rupee of cashback.")

    # A2 (Slide 7 / 14): Corrected Detailed Problem Statement
    a2 = prs.slides.add_slide(blank)
    set_bg(a2, BG_LIGHT)
    add_chrome(a2, "APPENDIX A2 : DETAILED PROBLEM", "Problem Details, Dated Market Sources & Synthetic Persona", "Appendix A2 (Slide 7 / 14)", True)
    p_cards = [
        ("1. Silent Repeat Churn (25% to 35% Yearly - ASSUMPTION)", "Merchants receive isolated payment alerts without customer-level retention tracking, losing repeat buyers before cashflow drops are noticed."),
        ("2. Dashboard Friction (<8% WAU - ASSUMPTION)", "Counter-operated kirana, darshini, and pharmacy shops lack time or training to write prompts, export CSVs, or configure CRM rules."),
        ("3. Dead Hours & Rush Prep (SYNTHETIC STORE)", "In our 60-day synthetic Indiranagar ledger, 3 to 5 PM scans fall 70% below peak while weekend morning rushes risk staple stockouts."),
    ]
    for i, (pt, pd) in enumerate(p_cards):
        cy = 1.40 + i * 1.78
        add_box(a2, 0.55, cy, 6.45, 1.60, WHITE, border_rgb=BORDER)
        add_text(a2, 0.75, cy + 0.14, 6.05, 0.34, pt, size=13.0, bold=True, color=NAVY)
        add_text(a2, 0.75, cy + 0.52, 6.05, 0.95, pd, size=11.5, color=MUTED)
    add_box(a2, 7.20, 1.40, 5.58, 3.60, WHITE, border_rgb=BORDER)
    a2.shapes.add_picture(c1, Inches(7.54), Inches(1.46), width=Inches(4.90))
    add_box(a2, 7.20, 5.15, 5.58, 1.75, NAVY_DARK, border_rgb=CYAN)
    add_text(
        a2, 7.38, 5.26, 5.22, 1.55,
        "VERIFIED PAYTM DISCLOSURES vs SYNTHETIC PERSONA\n* Paytm Scale: 1.57 Cr device merchants, Rs 7.1L Cr GMV (Q1 FY27, Jul 2026); ~7% loan pen. (Q4 FY26, May 2026).\n* Synthetic Persona: Ramesh Rao, Sri Krishna Darshini (917 scans/30d, Rs 1.75L GMV; 48 regulars silent 16-38d).",
        size=10.6, color=WHITE
    )
    set_notes(a2, "Appendix A2 provides full source citations for Paytm's Q1 FY27 and Q4 FY26 disclosures and explicitly distinguishes assumed SMB churn rates from our synthetic store dataset.")

    # A3 (Slide 8 / 14): Corrected 3-Pillar Solution Mapping
    a3 = prs.slides.add_slide(blank)
    set_bg(a3, BG_LIGHT)
    add_chrome(a3, "APPENDIX A3 : 3-PILLAR MAPPING", "Track 1 Coverage: Growth, Operations, and Customer Recovery", "Appendix A3 (Slide 8 / 14)", True)
    pillars_a3 = [
        ("PILLAR 1 : GROW BUSINESS", "Retention & Dead-Hour Copilot", [
            ("Churn Detection (Built)", "SQL flags regulars silent 16+ days (16-38d in synthetic DB)."),
            ("Dead-Hour Cohort (Built)", "Identifies 3-5 PM snack buyers (55 detected, 34 eligible)."),
            ("QR Cashback Pass (Planned)", "Production hook into Paytm QR coupon redemption rail."),
            ("Holdout Split (Built)", "Randomized 80/20 split for honest pilot measurement."),
        ]),
        ("PILLAR 2 : MANAGE OPERATIONS", "Rush Prep & Cashflow Signals", [
            ("Rush Alert (Simulated)", "Flags weekend morning demand spikes from scan history."),
            ("Supplier PO (Simulated)", "Drafts pre-filled WhatsApp restock order (+18L milk, +12kg batter)."),
            ("Cap Ledger (Stage 1)", "Tracks committed cashback passes separately from API credits."),
            ("Credit Signal (Planned)", "Feeds verified retention metrics to partner lending funnel."),
        ]),
        ("PILLAR 3 : SERVE CUSTOMERS", "Privacy-Safe Service Recovery", [
            ("Delay Tagging (Simulated)", "Flags peak-hour queue delays and low feedback tags."),
            ("Goodwill Pass (Simulated)", "1-tap bilingual apology + Rs 30 next-scan coffee voucher."),
            ("Zero PII Leakage (Built)", "Merchant sees cohort counts/first-name hints, never raw phone numbers."),
            ("DPDP & Cooldown (Built)", "Hard opt-in filter and 14-day per-payer message cooldown."),
        ]),
    ]
    for i, (ptag, phead, pbul) in enumerate(pillars_a3):
        px = 0.55 + i * 4.14
        add_box(a3, px, 1.40, 3.94, 5.45, WHITE, border_rgb=BORDER)
        add_text(a3, px + 0.18, 1.56, 3.58, 0.28, ptag, size=10.5, bold=True, color=NAVY)
        add_text(a3, px + 0.18, 1.86, 3.58, 0.36, phead, size=14.5, bold=True, color=INK)
        add_bullet_list(a3, px + 0.14, 2.32, 3.66, 4.35, pbul, font_pt=12.0, space_pt=10)
    set_notes(a3, "Appendix A3 maps each capability across Track 1's three mandates and labels whether each feature is built in SQL today, simulated in the MVP, or planned for production.")

    # A4 (Slide 9 / 14): Corrected End-to-End Workflow & Synthetic Disclosure
    a4 = prs.slides.add_slide(blank)
    set_bg(a4, BG_LIGHT)
    add_chrome(a4, "APPENDIX A4 : WORKFLOW & MVP SCOPE", "Closing-Time Workflow (Synthetic Dataset & Deterministic Simulation)", "Appendix A4 (Slide 9 / 14)", True)
    steps_a4 = [
        ("STEP 1 : SQL SCAN (BUILT)", "Queries 1,613 synthetic UPI scans across 240 payers. Detects 48 evening regulars silent 16 to 38 days (median 19d)."),
        ("STEP 2 : GUARDRAILS (BUILT)", "Excludes 5 unconsented (DPDP) & 3 on 14d cooldown -> 40 eligible. Checks Rs 25/Rs 150 (16.7% <= 22% floor) & Rs 300 cap."),
        ("STEP 3 : APPROVAL (BUILT)", "Plays draft Kannada/Hindi/English prompt via browser TTS; merchant approves by 1-tap button or voice command (32 nudged, 8 control)."),
        ("STEP 4 : 48H SIMULATOR", "Deterministic simulation: 12/32 vs 1/8 return (p = 0.24). Generates Rs 1,740 sales (5.8x sales/promo) or +Rs 222 pre-fee profit at 30% margin."),
    ]
    for i, (sh, sb) in enumerate(steps_a4):
        sx = 0.55 + i * 3.10
        add_box(a4, sx, 1.45, 2.92, 3.10, WHITE, border_rgb=NAVY, border_pt=1.2)
        add_text(a4, sx + 0.14, 1.60, 2.64, 0.36, sh, size=11.5, bold=True, color=NAVY)
        add_text(a4, sx + 0.14, 2.05, 2.64, 2.35, sb, size=11.5, color=INK)
    add_box(a4, 0.55, 4.80, 12.23, 2.05, AMB_SOFT, border_rgb=AMBER, border_pt=1.2)
    add_text(a4, 0.78, 4.95, 11.75, 0.34, "TRANSPARENCY DISCLOSURE: SYNTHETIC DATA & COMMITMENT ACCOUNTING", size=12.5, bold=True, color=AMBER)
    add_text(
        a4, 0.78, 5.34, 11.75, 1.38,
        "1. Dataset: All 240 customers and 1,613 historical UPI transactions are synthetically seeded (random.Random(20261025)); 48h redemptions are simulated.\n"
        "2. Cap Commitment Note: Issuing 32 passes of Rs 25 equals Rs 800 face exposure if capped only at scan time; strict commitment accounting caps active unredeemed passes at 12 per wave for a Rs 300 hard ceiling.",
        size=12.0, color=INK
    )
    set_notes(a4, "Appendix A4 replaces the old 'zero mock numbers' banner with an explicit disclosure of our synthetic seed and explains wave-based commitment accounting so merchant spend never exceeds Rs 300.")

    # A5 (Slide 10 / 14): Corrected Layered Tech Stack (Built vs Pending vs Planned)
    a5 = prs.slides.add_slide(blank)
    set_bg(a5, BG_LIGHT)
    add_chrome(a5, "APPENDIX A5 : SYSTEM ARCHITECTURE", "Component Status: Built in MVP vs Pending vs Production Plan", "Appendix A5 (Slide 10 / 14)", True)
    layers_a5 = [
        ("TOUCHPOINTS", "[BUILT] Companion PWA & Soundbox UI", "[BUILT] 1-Tap Button Fallback", "[PLANNED] Soundbox 4G Firmware & WhatsApp Cloud"),
        ("VOICE & LLM", "[BUILT] Browser Speech APIs (Draft Copy)", "[PENDING] Sarvam Bulbul/Saarika Adapter", "[OPTIONAL] Groq Free-Tier English Grounded Q&A"),
        ("CORE ENGINE", "[BUILT] Python FastAPI + SQL RFM Engine", "[BUILT] 5-Rule Deterministic Policy Engine", "[BUILT] 80/20 Holdout & Profit ROI Calculator"),
        ("DATA & OPS", "[BUILT] SQLite Ledger & Immutable Audit Log", "[PLANNED] Redis Atomic Redemption Counter", "[PLANNED] PostgreSQL & n8n Cron Workflows"),
    ]
    for idx, (lname, b1, b2, b3) in enumerate(layers_a5):
        ly = 1.42 + idx * 1.36
        add_box(a5, 0.55, ly, 12.23, 1.20, WHITE, border_rgb=BORDER)
        add_box(a5, 0.55, ly, 2.20, 1.20, NAVY_DARK, rounded=False)
        add_text(a5, 0.68, ly + 0.38, 1.95, 0.45, lname, size=11.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        for j, btxt in enumerate([b1, b2, b3]):
            bx = 2.92 + j * 3.25
            is_built = btxt.startswith("[BUILT]")
            is_pend = btxt.startswith("[PENDING]") or btxt.startswith("[OPTIONAL]")
            bg_c = EMER_SOFT if is_built else (AMB_SOFT if is_pend else SLATE_BG)
            bd_c = EMERALD if is_built else (AMBER if is_pend else BORDER)
            add_box(a5, bx, ly + 0.14, 3.10, 0.92, bg_c, border_rgb=bd_c, border_pt=1.1)
            add_text(a5, bx + 0.12, ly + 0.24, 2.86, 0.72, btxt, size=11.2, bold=True, color=INK)
    set_notes(a5, "Appendix A5 tags every single technology box as Built in MVP (green), Pending Sponsor Verification or Optional (amber), or Production Plan (grey). Unmeasured latency claims have been removed.")

    # A6 (Slide 11 / 14): Corrected USP Comparison Table
    a6 = prs.slides.add_slide(blank)
    set_bg(a6, BG_LIGHT)
    add_chrome(a6, "APPENDIX A6 : POSITIONING (USP)", "Differentiation vs Manual Khata Apps and Unconstrained Chatbots", "Appendix A6 (Slide 11 / 14)", True)
    tbl_s = a6.shapes.add_table(6, 4, Inches(0.55), Inches(1.42), Inches(12.23), Inches(5.30))
    tbl = tbl_s.table
    widths = [Inches(2.35), Inches(3.00), Inches(3.00), Inches(3.88)]
    for i, w in enumerate(widths):
        tbl.columns[i].width = w
    hdrs = ["Dimension", "Manual BI / Khata Apps", "Generic LLM Chatbots", "PaytmAI VyaparVoice"]
    rows_a6 = [
        ["Merchant Interaction", "Charts & manual ledger entry", "Requires typing business prompts", "Proactive 8:30 PM voice brief + 1-tap button"],
        ["Customer Graph", "Incomplete manual phonebooks", "No link to payment history", "Paytm QR repeat-payer IDs (synthetic in MVP)"],
        ["Financial Safety", "Manual coupon rules", "LLM can hallucinate unsafe offers", "Deterministic 5-rule cap, floor & DPDP filter"],
        ["Privacy (DPDP)", "Raw phone lists exposed", "Prompt data leakage risk", "Opt-in filter; zero raw PII sent to LLMs"],
        ["ROI Measurement", "Vanity message counts", "No control group", "80/20 holdout + profit ROI (simulated; pilot-ready)"],
    ]
    for c_i, h in enumerate(hdrs):
        cell = tbl.cell(0, c_i)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY_DARK if c_i < 3 else CYAN_DARK
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        p.text = assert_clean_text(h)
        p.font.name = FONT_MAIN
        p.font.size = Pt(12.0)
        p.font.bold = True
        p.font.color.rgb = WHITE
    for r_i, r_data in enumerate(rows_a6, start=1):
        for c_i, val in enumerate(r_data):
            cell = tbl.cell(r_i, c_i)
            cell.fill.solid()
            cell.fill.fore_color.rgb = EMER_SOFT if c_i == 3 else (WHITE if r_i % 2 == 1 else SLATE_BG)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.text = assert_clean_text(val)
            p.font.name = FONT_MAIN
            p.font.size = Pt(11.2)
            p.font.bold = c_i in (0, 3)
            p.font.color.rgb = INK
    set_notes(a6, "Appendix A6 contrasts VyaparVoice against manual BI dashboards and generic LLM chatbots, highlighting deterministic financial guardrails, DPDP privacy, and holdout-based profit accounting.")

    # A7 (Slide 12 / 14): Corrected Impact Targets & Holdout Details
    a7 = prs.slides.add_slide(blank)
    set_bg(a7, BG_LIGHT)
    add_chrome(a7, "APPENDIX A7 : PILOT TARGETS & STATS", "Hypothesized Pilot Targets vs Simulated MVP Metrics", "Appendix A7 (Slide 12 / 14)", True)
    add_box(a7, 0.55, 1.42, 6.50, 2.60, WHITE, border_rgb=BORDER)
    add_text(a7, 0.75, 1.54, 6.10, 0.30, "MERCHANT METRICS (SIMULATED MVP & PILOT TARGETS)", size=12.0, bold=True, color=EMERALD)
    add_bullet_list(
        a7, 0.72, 1.88, 6.15, 2.05,
        [
            ("Simulated Profit ROI", "+Rs 222 pre-fee (+74%) / +Rs 91.50 post-fee at 30% margin (5.8x sales/promo)."),
            ("Repeat Visit Target", "+14% to +22% monthly repeat visits (PILOT TARGET - unvalidated)."),
            ("Waste Reduction Target", "-18% afternoon perishable loss (PILOT TARGET - unvalidated)."),
        ],
        font_pt=11.5, space_pt=6
    )
    add_box(a7, 0.55, 4.20, 6.50, 2.65, WHITE, border_rgb=BORDER)
    add_text(a7, 0.75, 4.32, 6.10, 0.30, "PAYTM ECOSYSTEM HYPOTHESES (PILOT TARGETS)", size=12.0, bold=True, color=NAVY)
    add_bullet_list(
        a7, 0.72, 4.66, 6.15, 2.05,
        [
            ("Device Retention", "-35% Soundbox churn reduction (PILOT TARGET - unvalidated)."),
            ("Store QR Share", "+12% to +15% store Paytm QR GMV lift (PILOT TARGET - unvalidated)."),
            ("Credit Funnel", "Expand ~7% loan pen. (Q4 FY26) toward 15% (LONG-TERM TARGET)."),
        ],
        font_pt=11.5, space_pt=6
    )
    add_box(a7, 7.25, 1.42, 5.53, 3.70, WHITE, border_rgb=BORDER)
    a7.shapes.add_picture(c2, Inches(7.56), Inches(1.48), width=Inches(4.90))
    tiles_a7 = [
        ("1.74x", "Gross Profit / Promo\n(30% Margin Sim.)", EMERALD),
        ("p = 0.24", "Fisher Exact (N=40)\nNeeds ~46 / Arm", AMBER),
        ("Rs 5.00", "Cached AI Cost/Mo\n(Rs 28.80 Uncached)", NAVY),
    ]
    for i, (tv, tl, tc) in enumerate(tiles_a7):
        tx = 7.25 + i * 1.88
        add_box(a7, tx, 5.28, 1.76, 1.56, WHITE, border_rgb=tc, border_pt=1.3)
        add_text(a7, tx + 0.06, 5.44, 1.64, 0.48, tv, size=19, bold=True, color=tc, align=PP_ALIGN.CENTER)
        add_text(a7, tx + 0.06, 5.98, 1.64, 0.72, tl, size=9.8, bold=True, color=MUTED, align=PP_ALIGN.CENTER)
    set_notes(a7, "Appendix A7 replaces unvalidated impact claims with explicit Pilot Target labels and fixes the AI cost tile from under Rs 2 to the consistent Rs 5.00 cached figure, or Rs 28.80 uncached.")

    # A8 (Slide 13 / 14): Corrected Unit Economics & 3 Revenue Cases
    a8 = prs.slides.add_slide(blank)
    set_bg(a8, BG_LIGHT)
    add_chrome(a8, "APPENDIX A8 : UNIT ECONOMICS & ARR", "Sensitivity Across Revenue Streams and TTS Caching Rates", "Appendix A8 (Slide 13 / 14)", True)
    add_box(a8, 0.55, 1.42, 6.50, 5.42, WHITE, border_rgb=BORDER)
    add_bullet_list(
        a8, 0.74, 1.58, 6.12, 5.10,
        [
            ("Stream 1 (SaaS - Rs 79/mo)", "Bundled on Soundbox billing (~Rs 90-100/mo, FY25 AR) after 2 free trial campaigns -> Rs 74.4 Cr ARR at 5% base."),
            ("Stream 2 (Promo Fee - Rs 102/mo)", "5% fee on ~Rs 2,040/mo redeemed promo bills -> +Rs 96.1 Cr (Rs 170.5 Cr ARR combined with SaaS)."),
            ("Stream 3 (Credit Uplift - Rs 77/mo)", "1.5%-2.5% partner NBFC origination fee uplift -> +Rs 72.5 Cr (Rs 243.0 Cr Full Base ARR)."),
            ("Cached COGS (Rs 48.00/mo)", "Rs 1.80 SLM + Rs 3.20 TTS (~88% cache hit on Bulbul v3) + Rs 43 messaging = 81.4% gross margin on Rs 258 ARPU."),
            ("Uncached Sensitivity (Rs 71.80/mo)", "Uncached Bulbul v3 (9k chars/mo at Rs 30/10k = Rs 27.00 TTS) = 72.2% gross margin (Bulbul v2 Rs 13.50 = 77.4%)."),
        ],
        font_pt=11.5, space_pt=10
    )
    add_box(a8, 7.25, 1.42, 5.53, 3.85, WHITE, border_rgb=BORDER)
    a8.shapes.add_picture(c3, Inches(7.38), Inches(1.52), width=Inches(5.25))
    add_box(a8, 7.25, 5.42, 5.53, 1.42, NAVY_DARK, border_rgb=CYAN)
    add_text(
        a8, 7.42, 5.54, 5.18, 1.18,
        "CONSERVATIVE PRESENTATION ANCHOR\nLead with Rs 74.4 Cr (SaaS-only) to Rs 170.5 Cr (Core Product) ARR at 5% adoption; treat Rs 243.0 Cr (with credit referrals) as upside.",
        size=11.2, bold=True, color=WHITE
    )
    set_notes(a8, "Appendix A8 reconciles every rupee of ARPU, COGS, Sarvam Bulbul v3 versus v2 pricing, cache-hit assumptions, and the three ARR cases.")

    # A9 (Slide 14 / 14): Corrected Validation & Rollout Plan
    a9 = prs.slides.add_slide(blank)
    set_bg(a9, NAVY_DARK)
    add_box(a9, 0, 0, 13.333, 0.10, AMBER, rounded=False)
    add_text(a9, 0.65, 0.42, 6.5, 0.32, "APPENDIX A9 : VALIDATION & ROLLOUT ROADMAP", size=11.0, bold=True, color=CYAN)
    add_text(a9, 0.65, 0.85, 12.0, 0.50, "From Synthetic Hackathon Slice to Powered 50-Store Bengaluru Pilot", size=24, bold=True, color=WHITE)
    phases_a9 = [
        ("STAGE A : BUILT TODAY (MVP)", "Synthetic Slice & Guardrails", [
            "FastAPI + SQLite synthetic store (240 payers, 1,613 scans)",
            "5-rule policy engine + 80/20 holdout & profit ROI math",
            "Companion Soundbox UI + 21/21 automated QA checks",
        ]),
        ("STAGE B : DAYS 1 TO 45 PILOT", "Bengaluru 50-Store Validation", [
            "Verify Sarvam TTS/ASR in sponsor window & native copy review",
            "Run >=4 weekly waves (~50+ payers/arm) for 80% statistical power",
            "Measure real net merchant profit after cashback & 5% fee",
        ]),
        ("STAGE C : PRODUCTION SCALE", "Paytm Ecosystem Integration", [
            "Integrate Paytm QR dynamic coupon ledger & 4G Soundbox push",
            "Self-host open-weight SLMs (e.g., Apache 2.0 Sarvam-30B/M) if validated",
            "Roll out across Soundbox languages & partner credit funnel",
        ]),
    ]
    for i, (pt, ph, pitems) in enumerate(phases_a9):
        px = 0.65 + i * 4.10
        add_box(a9, px, 1.65, 3.85, 3.80, NAVY, border_rgb=CYAN, border_pt=1.2)
        add_text(a9, px + 0.18, 1.82, 3.50, 0.30, pt, size=10.5, bold=True, color=CYAN)
        add_text(a9, px + 0.18, 2.16, 3.50, 0.36, ph, size=14.5, bold=True, color=WHITE)
        for j, it in enumerate(pitems):
            add_text(a9, px + 0.18, 2.68 + j * 0.85, 3.50, 0.78, f"* {it}", size=11.2, color=RGBColor(0xE2, 0xE8, 0xF0))
    add_box(a9, 0.65, 5.75, 12.05, 1.10, RGBColor(0x00, 0x11, 0x33), border_rgb=CYAN)
    add_text(a9, 0.90, 5.92, 7.5, 0.36, "PaytmAI VyaparVoice | Team SemiTechnical (Kevin Joseph)", size=15.5, bold=True, color=WHITE)
    add_text(a9, 0.90, 6.32, 11.5, 0.34, "Deck v1 (Slide 14 / 14) | Live Demo: https://paytmai-vyaparvoice.onrender.com", size=11.5, color=CYAN)
    set_notes(a9, "Appendix A9 outlines the concrete validation steps from today's synthetic vertical slice to a statistically powered 50-store Bengaluru pilot and production integration.")

    prs.save(PPTX_OUT)
    # Also copy to /home/user/deck/v1/
    import shutil
    shutil.copy2(PPTX_OUT, os.path.join(OUT_DIRS[1], "PaytmAI_VyaparVoice_v1.pptx"))
    print("Saved Deck v1:", PPTX_OUT)


if __name__ == "__main__":
    build_deck_v1()
